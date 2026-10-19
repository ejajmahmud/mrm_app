"""Run after converting both PowerPoint files to PDF with LibreOffice.
Adds PDF navigation, exports PNG/SVG diagrams, checks text and geometry,
and prepares a download bundle. Dependencies: PyMuPDF, python-pptx, Pillow.
"""
from pathlib import Path
from pptx import Presentation
import fitz,json,re,zipfile,itertools
from PIL import Image,ImageDraw
ROOT=Path(__file__).resolve().parents[1]
NAMES=['01_Master_Platform_Mindmap','02_Muslim_User_Journeys','03_Care_and_Volunteer_Connections','04_Teacher_to_Learning_Connections','05_Islamic_Utility_Dependencies','06_Confidential_Support_and_Zakat','07_MRM_Operations_Connections','08_Feature_Connection_Index']
normalise=lambda s:re.sub(r'[^a-z0-9]','',s.lower())
results={}
for stem,count in [('MRM_Connected_Platform_58_Pages',58),('MRM_Feature_Connectivity_8_Maps',8)]:
 deck=Presentation(ROOT/(stem+'.pptx'));pdf=fitz.open(ROOT/(stem+'.pdf'));assert len(pdf)==len(deck.slides)==count
 # Existing PDF links are replaced deterministically, making this step repeatable.
 for page in pdf:
  for link in page.get_links():page.delete_link(link)
 ids={s.slide_id:i for i,s in enumerate(deck.slides)};linkcount=0;missing=[];outside=[]
 for i,(s,page) in enumerate(zip(deck.slides,pdf)):
  txt=normalise(page.get_text())
  for sh in s.shapes:
   if sh.has_text_frame and len(sh.text)>5 and normalise(sh.text) not in txt:missing.append((i+1,sh.text))
   if sh.left<0 or sh.top<0 or sh.left+sh.width>deck.slide_width+10 or sh.top+sh.height>deck.slide_height+10:outside.append((i+1,sh.name))
   dest=sh.click_action.target_slide
   if dest is not None:
    assert dest.slide_id in ids
    box=fitz.Rect(sh.left/12700,sh.top/12700,(sh.left+sh.width)/12700,(sh.top+sh.height)/12700)
    page.insert_link({'kind':fitz.LINK_GOTO,'from':box,'page':ids[dest.slide_id],'to':fitz.Point(0,0)});linkcount+=1
   # Verify extracted native text is inside the PDF page.
  for b in page.get_text('dict')['blocks']:
   if b['type']!=0:continue
   for ln in b['lines']:
    for sp in ln['spans']:
     x0,y0,x1,y1=sp['bbox']
     if x0<0 or y0<0 or x1>page.rect.width+.5 or y1>page.rect.height+.5:outside.append((i+1,sp['text']))
 assert not missing,missing
 assert not outside,outside
 titles=[]
 for i,s in enumerate(deck.slides):
  candidates=[sh for sh in s.shapes if sh.has_text_frame and .9*914400<sh.top<1.3*914400 and sh.width>8*914400]
  titles.append(candidates[0].text if candidates else f'Page {i+1}')
 pdf.set_toc([[1,t,i+1] for i,t in enumerate(titles)])
 temp=ROOT/(stem+'_linked.pdf');pdf.save(temp,garbage=4,deflate=True);pdf.close();temp.replace(ROOT/(stem+'.pdf'))
 check=fitz.open(ROOT/(stem+'.pdf'));assert sum(len(p.get_links()) for p in check)==linkcount
 results[stem]={'pages':count,'internal_links':linkcount,'text_missing':len(missing),'out_of_bounds':len(outside)}
# Check all graph nodes are disjoint; edges must not pass through unrelated nodes.
model=json.loads((ROOT/'Feature_Connectivity_Model.json').read_text());collisions=[];crossings=[]
def overlap(a,b):
 return min(a['x']+a['w'],b['x']+b['w'])>max(a['x'],b['x'])+.01 and min(a['y']+a['h'],b['y']+b['h'])>max(a['y'],b['y'])+.01
for m in model:
 for (ka,a),(kb,b) in itertools.combinations(m['nodes'].items(),2):
  if overlap(a,b):collisions.append((m['page'],ka,kb))
 for e in m['edges']:
  for aa,bb in zip(e['points'],e['points'][1:]):
   for k,n in m['nodes'].items():
    if k in [e['source'],e['target']]:continue
    for step in range(1,100):
     q=step/100;x=aa[0]+q*(bb[0]-aa[0]);y=aa[1]+q*(bb[1]-aa[1])
     if n['x']+.015<x<n['x']+n['w']-.015 and n['y']+.015<y<n['y']+n['h']-.015:
      crossings.append((m['page'],e['source'],e['target'],k));break
assert not collisions,collisions
assert not crossings,crossings
results['geometry']={'node_overlaps':len(collisions),'edges_through_unrelated_nodes':len(crossings),'maps':len(model),'nodes':sum(len(m['nodes']) for m in model),'relations':sum(len(m['edges']) for m in model)}
# Export full-resolution standalone vector and image diagrams.
(ROOT/'previews').mkdir(exist_ok=True);(ROOT/'vectors').mkdir(exist_ok=True)
pdf=fitz.open(ROOT/'MRM_Feature_Connectivity_8_Maps.pdf')
contact=Image.new('RGB',(1800,2200),'#111B35');draw=ImageDraw.Draw(contact)
for i,page in enumerate(pdf):
 pix=page.get_pixmap(matrix=fitz.Matrix(2560/page.rect.width,2560/page.rect.width));pix.save(ROOT/'previews'/(NAMES[i]+'.png'))
 (ROOT/'vectors'/(NAMES[i]+'.svg')).write_text(page.get_svg_image(text_as_path=True))
 im=Image.frombytes('RGB',[pix.width,pix.height],pix.samples);im.thumbnail((890,510));x=i%2*900;y=i//2*550;contact.paste(im,(x,y+25));draw.text((x+10,y+6),f'MAP {i+1:02d}',fill='white')
contact.save(ROOT/'previews'/'All_8_Maps_Overview.jpg')
poster=fitz.open();poster.insert_pdf(pdf,from_page=0,to_page=0);poster.save(ROOT/'MRM_Master_Feature_Mindmap.pdf',garbage=4,deflate=True)
results['exports']={'png_diagrams':8,'svg_diagrams':8,'master_poster_pages':1}
(ROOT/'Verification.json').write_text(json.dumps(results,indent=2))
files=['MRM_Connected_Platform_58_Pages.pptx','MRM_Connected_Platform_58_Pages.pdf','MRM_Feature_Connectivity_8_Maps.pptx','MRM_Feature_Connectivity_8_Maps.pdf','MRM_Master_Feature_Mindmap.pdf','Map_Contents.md','Feature_Connectivity_Model.json']
with zipfile.ZipFile(ROOT/'MRM_Connectivity_Download.zip','w',zipfile.ZIP_DEFLATED) as z:
 for f in files:z.write(ROOT/f,f)
 for folder in ['previews','vectors']:
  for f in sorted((ROOT/folder).iterdir()):z.write(f,f.relative_to(ROOT))
print(json.dumps(results,indent=2))
print('Download bundle:',round((ROOT/'MRM_Connectivity_Download.zip').stat().st_size/1024/1024,1),'MiB')
