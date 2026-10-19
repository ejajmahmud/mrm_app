"""Build the editable MRM connectivity chapter from the existing 50-page deck.
Dependencies: python-pptx, Pillow. Run from any directory with Python 3.
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
import json,re,io
from PIL import ImageFont
BOLD_FONT=ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",160)
def heading_size(title,width,maxsize=16):
 return min(maxsize,width*72/(BOLD_FONT.getlength(title)/160)*.94)

HERE=Path(__file__).resolve().parents[1]
SOURCE=HERE.parent/'MRM_50_Page_Deck.pptx'
NAVY='111B35'; PANEL='1B2846'; WHITE='F7F9FF'; MUTED='BBC8DF'; CYAN='62DAE8'; PURPLE='C2A3FF'; GOLD='E8C578'; GREEN='8DD8B1'; LINE='34435F'; FONT='DejaVu Sans'
COLORS=[CYAN,PURPLE,GREEN,GOLD]
original=Presentation(SOURCE)
assets=[original.slides[0].shapes[n].image.blob for n in [0,4,5]]
P=None;FULL=False;pending=[];maps=[];nodes={};edge_records=[]

def rgb(c):return RGBColor.from_string(c)
def box(s,x,y,w,h,fill=PANEL,line=None,alpha=None,round=True):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round else MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
 sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill)
 if line:sh.line.color.rgb=rgb(line);sh.line.width=Pt(1)
 else:sh.line.fill.background()
 if round:
  try:sh.adjustments[0]=.10
  except Exception:pass
 if alpha is not None:
  e=OxmlElement('a:alpha');e.set('val',str(alpha));sh.fill._xPr.solidFill.srgbClr.append(e)
 return sh

def text(s,x,y,w,h,txt,size=15,color=WHITE,bold=False,align=PP_ALIGN.LEFT):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.clear();tf.word_wrap=True
 tf.margin_left=tf.margin_top=tf.margin_right=tf.margin_bottom=0
 for i,line in enumerate(txt.split('\n')):
  pa=tf.paragraphs[0] if i==0 else tf.add_paragraph();pa.text=line;pa.font.name=FONT;pa.font.size=Pt(size);pa.font.bold=bold;pa.font.color.rgb=rgb(color);pa.alignment=align;pa.space_after=Pt(2)
 return sh

def chip(s,x,y,w,txt,col=CYAN,target=None):
 sh=box(s,x,y,w,.34,line=LINE);tx=text(s,x+.10,y+.05,w-.2,.26,txt,10,col,True)
 if target and FULL:pending.extend([(sh,target),(tx,target)])
 return sh

def base(title,subtitle,topic='FEATURE CONNECTIVITY'):
 global nodes,edge_records
 nodes={};edge_records=[]
 s=P.slides.add_slide(P.slide_layouts[6]);page=len(P.slides)
 s.shapes.add_picture(io.BytesIO(assets[0]),0,0,width=P.slide_width,height=P.slide_height)
 box(s,0,0,16,9,NAVY,alpha=90000,round=False)
 box(s,.7,.48,.045,.30,CYAN,round=False)
 text(s,.9,.49,11,.3,topic+'  /  MRM DIGITAL COMMUNITY PLATFORM',10,CYAN,True)
 s.shapes.add_picture(io.BytesIO(assets[1]),Inches(13.5),Inches(.38),height=Inches(.38))
 s.shapes.add_picture(io.BytesIO(assets[2]),Inches(14.45),Inches(.40),height=Inches(.32))
 text(s,.7,1.03,14.6,.66,title,31,WHITE,True)
 text(s,.72,1.8,14.5,.55,subtitle,14,MUTED)
 box(s,.7,8.3,14.6,.012,LINE,round=False)
 tx=text(s,.7,8.48,3.6,.25,'MRM  /  ZYLAX TECHNOLOGIES',9,MUTED)
 pending.append((tx,4 if FULL else 1))
 text(s,4.15,8.48,8.0,.26,'PROPOSED CONNECTIONS  /  EDITABLE NATIVE DIAGRAM',9,MUTED,align=PP_ALIGN.CENTER)
 text(s,14.0,8.43,1.3,.35,f'{page:02d} / {58 if FULL else 8}',12,CYAN,True,PP_ALIGN.RIGHT)
 maps.append({'page':page,'title':title,'nodes':nodes,'edges':edge_records})
 return s

def band(s,txt,col=CYAN):
 box(s,.7,7.8,14.6,.36)
 text(s,.9,7.86,14.15,.26,txt,10.5,col)

def node(s,key,x,y,w,h,title,body='',col=CYAN,target=None,code=None,hub=False):
 sh=box(s,x,y,w,h,fill='223653' if hub else PANEL,line=col)
 box(s,x+.02,y+.17,.035,h-.34,col,round=False)
 if code:
  circ=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x+.18),Inches(y+.15),Inches(.33),Inches(.33));circ.fill.solid();circ.fill.fore_color.rgb=rgb(NAVY);circ.line.color.rgb=rgb(col)
  text(s,x+.18,y+.205,.33,.19,code,8,col,True,PP_ALIGN.CENTER)
  tx=text(s,x+.65,y+.16,w-.85,.43,title,heading_size(title,w-.85,20 if hub else 16),WHITE,True)
 else:tx=text(s,x+.22,y+.17,w-.44,.48,title,heading_size(title,w-.44,20 if hub else 16),WHITE,True)
 if body:text(s,x+.23,y+.66,w-.46,h-.73,body,12 if not hub else 13,MUTED)
 if target and FULL:pending.extend([(sh,target),(tx,target)])
 nodes[key]={'x':x,'y':y,'w':w,'h':h,'title':title,'body':body,'detail_page':target,'color':col}
 return sh

def anchor(key,side):
 d=nodes[key];x,y,w,h=[d[k] for k in ['x','y','w','h']]
 return {'L':(x,y+h/2),'R':(x+w,y+h/2),'T':(x+w/2,y),'B':(x+w/2,y+h)}[side]

def line(s,points,col=CYAN,arrow=True,dashed=False):
 for i,(a,b) in enumerate(zip(points,points[1:])):
  sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(a[0]),Inches(a[1]),Inches(b[0]),Inches(b[1]));sh.line.color.rgb=rgb(col);sh.line.width=Pt(1.6)
  if arrow and i==len(points)-2:
   e=OxmlElement('a:tailEnd');e.set('type','triangle');e.set('w','sm');e.set('len','sm');sh.line._get_or_add_ln().append(e)
  if dashed:
   e=OxmlElement('a:prstDash');e.set('val','dash');sh.line._get_or_add_ln().append(e)
  # Keep edges beneath all nodes and labels, above the background.
  el=sh._element;el.getparent().remove(el);s.shapes._spTree.insert(4,el)

def label(s,x,y,w,txt,col=CYAN):
 if w<.72:
  x-=(.72-w)/2;w=.72
 box(s,x,y,w,.28,fill=NAVY,round=False)
 text(s,x+.02,y+.03,w-.04,.23,txt,9,col,True,PP_ALIGN.CENTER)

def edge(s,a,b,sa='R',sb='L',col=CYAN,via=None,txt=None,at=None,dashed=False,arrow=True):
 pts=[anchor(a,sa)]+(via or [])+[anchor(b,sb)]
 line(s,pts,col,arrow,dashed)
 if txt and at:label(s,*at,txt,col)
 edge_records.append({'source':a,'target':b,'meaning':txt or 'capability branch','points':pts,'dashed':dashed,'directional':arrow})

def finalize(s,summary):
 s.notes_slide.notes_text_frame.text=summary+'\n\nThis is a proposed feature-connectivity specification based on supplied scope documents and the current audience correction: all Muslims may use the platform. It is not proof of implemented behaviour. Marketplace excluded. Teacher subscriptions do not gate member assistance or volunteering. Religious sources, MRM policies, billing configuration and partner access require the agreed approvals. Solid directed arrows indicate a service relationship or handoff; dashed arrows indicate an explicit conditional/feedback path. These diagrams do not imply unrestricted data sharing.'


def build_chapter():
 # 01 / 51 -- broad map; all feature families with distinct privacy-aware branches.
 s=base('The complete platform, connected','A mind map of daily practice, learning, community care and the services that make them work together.')
 node(s,'hub',6.22,4.35,3.56,1.48,'MRM PLATFORM','For all Muslims\nShared identity • safe access',CYAN,hub=True)
 groups=[('faith',.7,2.58,'Daily Islamic practice','Prayer • Qibla • mosques • Quran\nDua • dhikr • hadith • Hijri / Ramadan\nInformational zakat estimate',GOLD,32,'01'),('learn',10.50,2.58,'Learning & teachers','Teacher approval • posts • teaching ads\nClasses • recordings • events\nTeacher-only publishing subscriptions',PURPLE,24,'02'),('care',.7,4.44,'Care & volunteering','Help requests • matching • assignments\nVerification • interview • gender rules\nNearby discovery • badges / ranks',GREEN,13,'03'),('private',10.50,4.44,'Confidential support & zakat','Restricted faith / family cases\nConsent • referral • partner follow-up\nAuthority decides zakat outcomes',GREEN,18,'04'),('public',.7,6.20,'Website & sponsors','Public information • participation routes\nApproved resources • sponsor recognition\nScheduled app / website placements',CYAN,31,'05'),('ops',10.50,6.20,'MRM operations','Accounts • moderation • configurations\nAlerts • activity • audit • reports\nPrivacy • backups • restricted exports',CYAN,42,'06')]
 for key,x,y,title,body,col,target,code in groups:node(s,key,x,y,4.8,1.48,title,body,col,target,code)
 # Compact three-arm trunks make the branches visually legible.
 for key in ['faith','care','public']:
  end=anchor(key,'R');start=anchor('hub','L');edge(s,'hub',key,'L','R',nodes[key]['color'],via=[(5.85,start[1]),(5.85,end[1])],arrow=False)
 for key in ['learn','private','ops']:
  end=anchor(key,'L');start=anchor('hub','R');edge(s,'hub',key,'R','L',nodes[key]['color'],via=[(10.15,start[1]),(10.15,end[1])],arrow=False)
 band(s,'Select a branch to visit its feature pages. Detailed maps label approvals, handoffs, access and feedback.' if FULL else 'Companion to the 58-page deck. Branch colours group features; detailed maps explain approvals and handoffs.')
 finalize(s,'Master mind map. Each branch is a capability family rather than an unrestricted data path. Sifat Names, Haram Codes and donation/payment administration remain supplied template functions with the scope qualifications recorded on pages 38–41.')
 # 02 / 52 -- customer-facing / member / roles mind map.
 s=base('One Muslim user, many ways to participate','Daily tools welcome the wider community; personal services and privileged roles use the appropriate access checks.')
 node(s,'person',6.2,4.25,3.6,1.65,'ALL MUSLIM USERS','Daily value • community care\nAccount for personal services',CYAN,hub=True)
 for key,x,y,title,body,col,target,code in [
  ('practice',.7,2.60,'Practise daily','Prayer, Quran and remembrance\nLocation choices for relevant tools',GOLD,32,'01'),
  ('learning',10.50,2.60,'Discover & learn','Approved programmes and teachers\nRegister where required; revisit recordings',PURPLE,24,'02'),
  ('help',.7,4.45,'Seek support','Private request, safe contact preference\nTrack only your authorised case details',GREEN,13,'03'),
  ('volunteer',10.50,4.45,'Contribute as a volunteer','Evidence + interview + MRM approval\nEligible offers and assigned service',GREEN,15,'04'),
  ('updates',.7,6.20,'Follow personal updates','Application / request status and reminders\nSafe alerts open the authorised detail',CYAN,11,'05'),
  ('teacher',10.50,6.20,'Teach through the platform','Teacher approval + valid plan entitlement\nSubmit educational content for review',PURPLE,25,'06')]:node(s,key,x,y,4.8,1.48,title,body,col,target,code)
 for key in ['practice','help','updates']:
  a=anchor('person','L');b=anchor(key,'R');edge(s,'person',key,'L','R',nodes[key]['color'],via=[(5.85,a[1]),(5.85,b[1])],arrow=False)
 for key in ['learning','volunteer','teacher']:
  a=anchor('person','R');b=anchor(key,'L');edge(s,'person',key,'R','L',nodes[key]['color'],via=[(10.15,a[1]),(10.15,b[1])],arrow=False)
 band(s,'Core member assistance and volunteer participation remain free. Teacher subscription privileges stay separate.')
 finalize(s,'Member and role participation map. Account identity, volunteer approval, teacher approval and subscription entitlement are separate controls. Public website discovery connects users to app entry and relevant participation routes.')
 # 03 / 53 -- care grid.
 s=base('How care, volunteering and recognition connect','Routine assignments and restricted concerns follow distinct access paths, with MRM coordinating the outcome.')
 xs=[.7,4.55,8.4,12.25];ys=[2.66,4.47,6.28];w=3.05;h=1.23
 data=[
 ('request',0,0,'Member request','Need, urgency and safe contact',GREEN,13),('triage',1,0,'MRM triage','Review sensitivity and suitability',CYAN,14),('sensitive',2,0,'Restricted concern','Private faith / family intake',GOLD,18),('staff',3,0,'Designated personnel','Controlled advice and follow-up',GOLD,19),
 ('approved',0,1,'Approved volunteer','Active status and permitted service',GREEN,15),('match',1,1,'Eligible matching','Area, category, language, gender',CYAN,14),('assign',2,1,'Confirmed assignment','Seeker / coordinator confirmation',CYAN,17),('close',3,1,'Progress & closure','Verified outcome; escalation if needed',GREEN,17),
 ('apply',0,2,'Evidence & interview','Evidence and MRM decision',GREEN,15),('oversee',1,2,'Coordinator oversight','Reassign or follow up unanswered work',CYAN,18),('activity',2,2,'Alerts & activity','Only authorised status / action',CYAN,23),('rank',3,2,'Service recognition','MRM criteria; correction / appeal',GOLD,21)]
 for key,c,r,a,b,col,t in data:node(s,key,xs[c],ys[r],w,h,a,b,col,t)
 edge(s,'request','triage',txt='submit',at=(4.04,3.00,.4))
 edge(s,'triage','sensitive',col=GOLD,txt='private',at=(7.80,3.00,.4))
 edge(s,'sensitive','staff',col=GOLD,txt='assign',at=(11.56,3.00,.38))
 edge(s,'triage','match','B','T',txt='routine',at=(5.43,4.02,1.0))
 edge(s,'apply','approved','T','B',col=GREEN,txt='verify',at=(1.86,5.94,1.0))
 edge(s,'approved','match',col=GREEN,txt='eligible',at=(4.04,4.82,.4))
 edge(s,'match','assign',txt='offer',at=(7.80,4.82,.4))
 edge(s,'assign','close',col=GREEN,txt='serve',at=(11.56,4.82,.38))
 edge(s,'match','oversee','B','T',txt='no response',at=(5.20,5.96,1.4),dashed=True)
 edge(s,'assign','activity','B','T',txt='updates',at=(9.36,5.96,1.0))
 edge(s,'close','rank','B','T',col=GOLD,txt='verify credit',at=(12.81,5.96,1.65))
 band(s,'Approval gates access; assignment gates case detail. Rank never grants sensitive-case or administrator access.')
 finalize(s,'Care network. MRM triage routes routine requests to eligible matching, while restricted faith/family matters go to designated personnel. Volunteer evidence review and online interview precede approval. Confirmation precedes assigned details. Unanswered work returns to coordinator oversight. Service credit requires verified outcomes.')
 # 04 / 54 -- teacher serpentine graph.
 s=base('From approved teacher to trusted learning','Teacher plans enable publishing allowances; MRM moderation controls what learners can actually see.')
 data=[
 ('apply',0,0,'Teacher application','Credentials, subjects and evidence',PURPLE,25),('approval',1,0,'MRM teacher approval','Authorised professional role',PURPLE,25),('plan',2,0,'Select teacher plan','MRM terms; applicable store billing',PURPLE,27),('entitlement',3,0,'Verified entitlement','Active paid / configured free plan',CYAN,28),
 ('published',0,1,'Approved publication','Posts, teaching ads and programmes',GREEN,26),('moderation',1,1,'MRM moderation','Approve, return, schedule or withdraw',GREEN,26),('submission',2,1,'Teacher submission','Content, class proposal or recording',PURPLE,26),('allowance',3,1,'Publishing allowance','Apply active plan limits',PURPLE,27),
 ('discover',0,2,'Learner discovery','Topic, language and teacher',CYAN,24),('details',1,2,'Programme & access','Details and registration rules',CYAN,29),('attend',2,2,'Live class / event','Reminder and permitted session access',CYAN,29),('revisit',3,2,'Approved recordings','Continue learning after the session',CYAN,24)]
 for key,c,r,a,b,col,t in data:node(s,key,xs[c],ys[r],w,h,a,b,col,t)
 for a,b in [('apply','approval'),('approval','plan'),('plan','entitlement')]:edge(s,a,b,col=PURPLE)
 edge(s,'entitlement','allowance','B','T',col=PURPLE,txt='enables',at=(13.04,4.02,1.0))
 for a,b in [('allowance','submission'),('submission','moderation'),('moderation','published')]:edge(s,a,b,'L','R',col=PURPLE)
 edge(s,'published','discover','B','T',col=GREEN,txt='publish',at=(1.86,5.96,1.0))
 for a,b in [('discover','details'),('details','attend'),('attend','revisit')]:edge(s,a,b,col=CYAN)
 # Return-for-revision feedback stays in the dedicated row gap.
 edge(s,'moderation','submission','B','B',col=GOLD,via=[(6.075,6.03),(9.925,6.03)],txt='return for revision',at=(6.9,5.89,2.1),dashed=True)
 band(s,'Payment does not approve content or grant staff access. Learners and volunteers do not need teacher subscriptions.')
 finalize(s,'Teacher-to-learner connectivity. Teacher approval precedes plan selection. Verified entitlements gate publishing allowances, while moderation independently gates publication. Revision is a feedback path. Expiry pauses paid privileges and preserves history; confirm historical-content visibility. Live sessions may use approved external meeting links; no custom conferencing infrastructure is implied.')
 # 05 / 55 -- inputs and feature outputs.
 s=base('Daily utilities connect through approved inputs','Location, religious content and local conventions feed the right tools without mixing unrelated permissions.')
 for x,title in [(.7,'INPUT / USER CHOICE'),(5.10,'VALIDATION / CONFIGURATION'),(10.20,'CONNECTED EXPERIENCE')]:text(s,x,2.35,4.7,.25,title,10,CYAN,True)
 for key,x,y,ww,a,b,col,t in [
 ('location',.7,2.87,3.7,'Location or manual area','Consent and device capability',CYAN,12),('local',5.1,2.87,4.3,'Local prayer & map settings','Local times, direction and place sources',CYAN,33),('practice',10.2,2.87,5.1,'Prayer • Qibla • nearby mosques','Useful daily guidance; permission handling',GOLD,34),
 ('sources',.7,4.59,3.7,'Religious sources & rights','Approved references and translations',PURPLE,41),('content',5.1,4.59,4.3,'Reviewed content library','Quran text / recitation and utility content',PURPLE,35),('remembrance',10.2,4.59,5.1,'Quran • dua • dhikr • hadith','Connect practice to approved learning',PURPLE,36),
 ('conventions',.7,6.31,3.7,'Calendar / zakat sources','MRM-approved local conventions',GREEN,39),('configuration',5.1,6.31,4.3,'Display & calculation rules','Source review; local adjustment policy',GREEN,40),('seasonal',10.2,6.31,5.1,'Hijri • Ramadan • zakat estimate','Guidance only; referral is a separate service',GREEN,39)]:node(s,key,x,y,ww,1.20,a,b,col,t)
 for a,b,c in [('location','local',CYAN),('local','practice',CYAN),('sources','content',PURPLE),('content','remembrance',PURPLE),('conventions','configuration',GREEN),('configuration','seasonal',GREEN)]:edge(s,a,b,col=c)
 band(s,'Daily content ↔ approved learning. Zakat estimate ↔ referral information; eligibility stays with the authority.')
 finalize(s,'Utility dependencies. Location choices affect location-related tools, not blanket access to personal information. Religious references, translations and recitation rights are reviewed separately. Calendar conventions and calculator sources require approval. Utility-to-learning links reference approved content. A zakat estimate does not automate assistance eligibility or disbursement.')
 # 06 / 56 -- privacy two lanes.
 s=base('Confidential support and zakat: controlled handoffs','These journeys share access controls and case accountability, not unrestricted volunteer or teacher visibility.')
 chip(s,.7,2.35,3.7,'CONSENT + MINIMUM NECESSARY DATA',GOLD)
 chip(s,4.62,2.35,4.25,'ROLE + RECORD-LEVEL AUTHORISATION',CYAN)
 chip(s,9.1,2.35,6.2,'SAFE ALERTS + RESTRICTED EVIDENCE + AUDIT',GREEN)
 text(s,.73,2.94,14,.3,'CONFIDENTIAL FAITH / FAMILY SUPPORT',11,GOLD,True)
 text(s,.73,5.10,14,.3,'ZAKAT ENQUIRY / REFERRAL',11,GREEN,True)
 for key,x,y,ww,a,b,col,t in [
 ('intake',.7,3.42,3.7,'Private concern','Safe contact and permitted information',GOLD,18),('restricted',5.1,3.42,4.3,'Designated MRM review','Private queue and appropriate assignment',GOLD,18),('adviser',10.2,3.42,5.1,'Authorised guidance & follow-up','Controlled closure or permitted referral',GOLD,19),
 ('enquiry',.7,5.58,3.7,'Enquiry + referral consent','Need, region and safe contact preference',GREEN,20),('review',5.1,5.58,4.3,'MRM completeness review','Approved forms, sharing method and contact',GREEN,20),('authority',10.2,5.58,5.1,'Zakat authority handoff','Authority decides eligibility and disbursement',GREEN,20)]:node(s,key,x,y,ww,1.30,a,b,col,t)
 edge(s,'intake','restricted',col=GOLD);edge(s,'restricted','adviser',col=GOLD)
 edge(s,'enquiry','review',col=GREEN);edge(s,'review','authority',col=GREEN)
 text(s,.75,7.16,14.5,.42,'Authorised status → case history → permitted user update → aggregate reporting. Partner API access is conditional.',14,MUTED)
 band(s,'No automated religious judgement. No guaranteed partner API. No zakat collection or disbursement in the proposal.')
 finalize(s,'Two restricted service lanes. General volunteers and teacher subscriptions do not grant access to confidential cases. Zakat referrals use consent and approved data sharing. Status feedback is recorded only where permitted. Relevant authority decisions are external; direct integration depends on authorisation, supported API, documentation and test access.')
 # 07 / 57 -- operations-centred mind map with specific relations.
 s=base('MRM operations connect content, revenue and oversight','Staff actions create controlled outputs across the app and website; service records feed aggregate reporting.')
 node(s,'operations',6.22,4.35,3.56,1.48,'MRM OPERATIONS','Duty-based staff access\nConfigure • review • record',CYAN,42,hub=True)
 groups=[
 ('public',.7,2.58,'Public website & discovery','Approved pages, resources and events\nVolunteer / teacher / support entry routes',CYAN,31,'01'),
 ('sponsor',10.5,2.58,'Sponsors & placements','Approved assets + start / end dates\nApp / website recognition; aggregate metrics',GOLD,30,'02'),
 ('publish',.7,4.44,'Content & learning publication','Teacher proposals and utility references\nModeration controls visible material',PURPLE,26,'03'),
 ('billing',10.5,4.44,'Teacher plans & entitlements','Client-owned store billing configuration\nVerified plan state → publishing allowances',PURPLE,27,'04'),
 ('services',.7,6.20,'Alerts • activity • audit','Service event → permitted recipient\nSafe alert → authorised detail / next action',CYAN,23,'05'),
 ('reports',10.5,6.20,'Reports & service oversight','Case / volunteer / referral status\nAgreed programme and sponsor summaries',GREEN,46,'06')]
 for key,x,y,a,b,col,t,code in groups:node(s,key,x,y,4.8,1.48,a,b,col,t,code)
 for key in ['public','publish','services']:
  a=anchor('operations','L');b=anchor(key,'R');edge(s,'operations',key,'L','R',nodes[key]['color'],via=[(5.85,a[1]),(5.85,b[1])],arrow=True)
 for key in ['sponsor','billing']:
  a=anchor('operations','R');b=anchor(key,'L');edge(s,'operations',key,'R','L',nodes[key]['color'],via=[(10.15,a[1]),(10.15,b[1])],arrow=True)
 a=anchor('reports','L');b=anchor('operations','R');edge(s,'reports','operations','L','R',GREEN,via=[(10.15,a[1]),(10.15,b[1])],arrow=True)
 band(s,'Sponsors cannot access private cases. Sponsorship never changes assignment, confidential guidance or zakat decisions.')
 finalize(s,'Operations map. MRM supplies governance, content approval, teacher plan rules and sponsor campaign inputs. Staff publishing feeds app/website experiences; billing state enables teacher entitlements, not content approval or staff privileges. Safe event delivery and audit support accountability. Aggregate reports inform staff oversight; identifiable exports remain restricted.')
 # 08 / 58 -- exact index / hyperlink guide.
 s=base('Connection index: what links to what','Use this index to trace a feature dependency back to the detailed scope pages in the full presentation.')
 rows=[
 ('Accounts & roles','Personal records / staff tools','Authorises appropriate access','Role + record + gender / sensitivity',43,CYAN),
 ('Volunteer verification','Discovery / request matching','Enables eligible service','Approved and active only',15,GREEN),
 ('Triage & matching','Assignment / case progress','Offers, confirms and follows up','Sensitive route stays restricted',13,GREEN),
 ('Teacher approval + plan','Posts / ads / programme proposals','Enables publishing allowances','Valid entitlement; no staff access',27,PURPLE),
 ('MRM moderation','Learning / website / utility content','Approves visible material','Religious review and content rights',26,PURPLE),
 ('Service events','Alerts / activity / case detail','Prompts the permitted next action','Privacy-safe alert; authorise detail',23,CYAN),
 ('Zakat enquiry','Authority referral / follow-up','Coordinates approved handoff','Consent; authority decides outcome',20,GREEN),
 ('Verified service outcome','Badges / ranks / service history','Awards evidence-based credit','MRM criteria; access unchanged',21,GOLD),
 ('Sponsor configuration','App / website placements','Schedules recognition / expiry','Approved assets; aggregate metrics',30,GOLD),
 ('Recorded outcomes','Operational reports / staff review','Aggregates service information','Restricted identifiable exports',46,CYAN)]
 x=[.7,3.35,6.55,10.05,14.55];widths=[2.65,3.2,3.5,4.5,.75];y=2.58;rh=.45
 box(s,.7,y,14.6,.50,fill='263B58',line=LINE)
 for j,hdr in enumerate(['FEATURE','CONNECTS TO','CONNECTION PURPOSE','ACCESS / POLICY CHECK','PAGE']):text(s,x[j]+.10,y+.14,widths[j]-.2,.25,hdr,9.5,CYAN,True)
 for i,(a,b,c,d,t,col) in enumerate(rows):
  yy=y+.50+i*rh;box(s,.7,yy,14.6,rh,fill=PANEL if i%2==0 else '17233D',round=False)
  box(s,.7,yy,.035,rh,col,round=False)
  for j,txt in enumerate([a,b,c,d,f'{t:02d}']):
   sh=text(s,x[j]+.10,yy+.12,widths[j]-.20,rh-.12,txt,10.5,col if j==4 else (WHITE if j==0 else MUTED),j in [0,4])
   if FULL:pending.append((sh,t))
 band(s,'Solid arrows = directed relationship / handoff. Dashed arrows = conditional path or revision feedback.')
 finalize(s,'Dependency index with ten key cross-feature relationships. Each row links to the corresponding detailed original deck page. The chapter maps the proposed product specification and does not certify deployment or complete implementation.')


def apply_links():
 for sh,t in pending:
  sh.click_action.target_slide=P.slides[t-1]

# Full 58-page presentation; original features and detailed page references remain valid.
P=Presentation(SOURCE);FULL=True;pending=[];maps=[]
for i,s in enumerate(P.slides,1):
 for sh in s.shapes:
  if sh.has_text_frame and re.fullmatch(r'\d{2} / 50',sh.text.strip()):
   for p in sh.text_frame.paragraphs:
    for r in p.runs:r.text=r.text.replace('/ 50','/ 58')
# Add a clear entry point on the original contents page, replacing its old helper band.
s=P.slides[3]
box(s,.7,7.77,14.6,.40,fill='253854',line=CYAN)
tx=text(s,.9,7.86,14,.25,'NEW  07  FEATURE CONNECTIVITY MIND MAPS   /   PAGES 51–58   /   SELECT TO EXPLORE',11,CYAN,True)
pending.extend([(tx,51),(list(s.shapes)[-2],51)])
build_chapter();apply_links();P.save(HERE/'MRM_Connected_Platform_58_Pages.pptx')
(HERE/'Feature_Connectivity_Model.json').write_text(json.dumps(maps,indent=2))
# Standalone chapter: all native nodes and connectors remain editable.
P=Presentation();P.slide_width=Inches(16);P.slide_height=Inches(9);FULL=False;pending=[];maps=[]
build_chapter();apply_links();P.save(HERE/'MRM_Feature_Connectivity_8_Maps.pptx')
(HERE/'Map_Contents.md').write_text('# MRM feature-connectivity maps\n\nAudience: all Muslims, with dedicated MRM revert-support services.\n\n'+ '\n'.join(f"{i+1:02d}. {m['title']} (full deck page {i+51})" for i,m in enumerate(maps))+'\n\nAll maps use editable PowerPoint shapes and connectors. They describe proposed feature relationships, with explicit privacy, approval and entitlement boundaries. The original 50-page deck is retained unchanged in the parent directory; the integrated copy has 58 pages.\n')
print('Built 58-page integrated deck and 8-page standalone chapter.')
