from pathlib import Path
import json,re,zipfile,hashlib
import fitz
from PIL import Image,ImageDraw
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[1]
N=lambda t:re.sub(r'[^a-z0-9]','',t.lower())
intro=fitz.open(ROOT/'MRM_User_Value_Bridges_24_Pages.pdf');original=fitz.open(ROOT.parent/'MRM_Strategic_Platform_72_Pages.pdf')
assert len(intro)==24 and len(original)==72
full=fitz.open();full.insert_pdf(intro);full.insert_pdf(original)
full.save(ROOT/'MRM_Complete_Connected_Platform_96_Pages.pdf',garbage=4,deflate=True);full.close()
checks={}
for stem,count in [('MRM_User_Value_Bridges_24_Pages',24),('MRM_Complete_Connected_Platform_96_Pages',96)]:
 deck=Presentation(ROOT/(stem+'.pptx'));pdf=fitz.open(ROOT/(stem+'.pdf'));assert len(deck.slides)==len(pdf)==count
 ids={s.slide_id:i for i,s in enumerate(deck.slides)};missing=[];outside=[];links=0
 for pg in pdf:
  for link in pg.get_links():pg.delete_link(link)
 for i,(sl,pg) in enumerate(zip(deck.slides,pdf)):
  content=N(pg.get_text())
  for sh in sl.shapes:
   if sh.has_text_frame and N(sh.text) and N(sh.text) not in content:missing.append([i+1,sh.text])
   if sh.left<0 or sh.top<0 or sh.left+sh.width>deck.slide_width+10 or sh.top+sh.height>deck.slide_height+10:outside.append([i+1,sh.name])
   # The new section must keep all text inside its own editable textbox.
   if i<24 and sh.has_text_frame and N(sh.text):
    r=fitz.Rect(sh.left/12700-1,sh.top/12700-1,(sh.left+sh.width)/12700+1,(sh.top+sh.height)/12700+1)
    if N(sh.text) not in N(pg.get_textbox(r)):outside.append([i+1,'Text outside parent',sh.text])
   target=sh.click_action.target_slide
   if target is not None:
    assert target.slide_id in ids
    r=fitz.Rect(sh.left/12700,sh.top/12700,(sh.left+sh.width)/12700,(sh.top+sh.height)/12700)
    pg.insert_link({'kind':fitz.LINK_GOTO,'from':r,'page':ids[target.slide_id],'to':fitz.Point(0,0)});links+=1
 assert not missing,missing
 assert not outside,outside
 titles=json.loads((ROOT/'Contents.json').read_text())
 toc=[[1,t,i+1] for i,t in enumerate(titles)]
 if count==96:toc += [[1,t,i+25] for i,t in enumerate(json.loads((ROOT.parent/'Main_Contents.json').read_text()))]
 pdf.set_toc(toc);temp=ROOT/(stem+'_linked.pdf');pdf.save(temp,garbage=4,deflate=True);pdf.close();temp.replace(ROOT/(stem+'.pdf'))
 check=fitz.open(ROOT/(stem+'.pdf'));assert sum(len(p.get_links()) for p in check)==links
 checks[stem]={'pages':count,'internal_links':links,'missing_text':0,'out_of_bounds':0,'new_section_textbox_overflows':0}
# Verify the original 72-page explanation is unchanged in the combined PDF.
combined=fitz.open(ROOT/'MRM_Complete_Connected_Platform_96_Pages.pdf')
assert all(original[i].get_text()==combined[i+24].get_text() for i in range(72))
# Each original photo remains a single media item; no screenshot duplication was introduced.
with zipfile.ZipFile(ROOT/'MRM_Complete_Connected_Platform_96_Pages.pptx') as z:
 media=[z.read(n) for n in z.namelist() if n.startswith('ppt/media/')]
 assert len(media)==len({hashlib.sha256(b).hexdigest() for b in media})==48
checks['preserved_original']={'pages_unchanged':72,'unique_media_files':48,'added_screenshot_uses':0}
checks['product_model']={'proposed_bridges':15,'required_fields':['person','need','yes','no','steps','carry','separate','benefit','signal'],'status':'Design proposals, not verified MVP behaviour'}
(ROOT/'Verification.json').write_text(json.dumps(checks,indent=2))
preview=ROOT/'previews';preview.mkdir(exist_ok=True)
doc=fitz.open(ROOT/'MRM_User_Value_Bridges_24_Pages.pdf')
for start in range(0,24,6):
 sheet=Image.new('RGB',(1800,1650),'#111B35');draw=ImageDraw.Draw(sheet)
 for j in range(start,start+6):
  pix=doc[j].get_pixmap(matrix=fitz.Matrix(.75,.75));im=Image.frombytes('RGB',[pix.width,pix.height],pix.samples);im.thumbnail((890,501));x=(j-start)%2*900;y=(j-start)//2*550;sheet.paste(im,(x,y+25));draw.text((x+10,y+5),f'PAGE {j+1:02d}',fill='white')
 sheet.save(preview/f'Pages_{start+1:02d}_{start+6:02d}.jpg')
for i,name in [(1,'Teacher_Example'),(2,'Connected_User_Journeys'),(4,'Teacher_to_Volunteer'),(5,'Volunteer_to_Training'),(6,'Seeker_to_Learning'),(19,'Teacher_Journey')]:
 doc[i].get_pixmap(matrix=fitz.Matrix(1.3334,1.3334)).save(preview/(name+'.png'))
print(json.dumps(checks,indent=2))
