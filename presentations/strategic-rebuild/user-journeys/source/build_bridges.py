"""Build the proposed user-value connection section with editable native shapes."""
from pathlib import Path
import io,json,math
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from PIL import ImageFont
from bridge_content import BRIDGES,BY_KEY
ROOT=Path(__file__).resolve().parents[1];OLD=ROOT.parent.parent/'MRM_50_Page_Deck.pptx'
ref=Presentation(OLD);ASSETS=[ref.slides[0].shapes[n].image.blob for n in [0,4,5]]
P=Presentation();P.slide_width=Inches(16);P.slide_height=Inches(9)
NAVY='111B35';PANEL='1B2846';LINE='34435F';WHITE='F7F9FF';MUTED='BBC8DF';CYAN='62DAE8';GOLD='E8C578';GREEN='8DD8B1';PURPLE='C2A3FF'
LINKS=[];TITLES=[];TEXTCHECK=[]
RF=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',200)
BF=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf',200)
def rgb(s):return RGBColor.from_string(s)
def box(s,x,y,w,h,c=PANEL,line=LINE):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=rgb(c);sh.line.color.rgb=rgb(line);sh.line.width=Pt(.8);sh.adjustments[0]=.08;return sh

def wrap(text,w,size,bold=False):
 f=BF if bold else RF;limit=w*72*200/size;out=[]
 for para in text.split('\n'):
  cur=''
  for word in para.split():
   if cur and f.getlength(cur+' '+word)>limit:out.append(cur);cur=word
   else:cur=(cur+' '+word).strip()
  out.append(cur)
 return out

def txt(s,x,y,w,h,text,size=14,c=WHITE,bold=False,align=PP_ALIGN.LEFT,target=None):
 # Wrap explicitly using actual font metrics so exported lines remain inside their bounds.
 lines=wrap(text,w,size,bold)
 while len(lines)*size*1.23>h*72 and size>9:
  size-=.25;lines=wrap(text,w,size,bold)
 assert len(lines)*size*1.23<=h*72,(len(P.slides),text,h,size)
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.word_wrap=False;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 for i,line in enumerate(lines):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='DejaVu Sans';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(c);p.alignment=align;p.space_after=Pt(0);p.line_spacing=1.1
 TEXTCHECK.append(dict(page=len(P.slides),text=text,x=x,y=y,w=w,h=h))
 if target:LINKS.append((sh,target))
 return sh

def arrow(s,x1,y1,x2,y2,c=CYAN):
 sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));sh.line.color.rgb=rgb(c);sh.line.width=Pt(1.7)
 e=OxmlElement('a:tailEnd');e.set('type','triangle');sh.line._get_or_add_ln().append(e)

def base(title,sub,section='USER-VALUE CONNECTIONS'):
 s=P.slides.add_slide(P.slide_layouts[6]);TITLES.append(title)
 s.shapes.add_picture(io.BytesIO(ASSETS[0]),0,0,width=P.slide_width,height=P.slide_height)
 sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,P.slide_width,P.slide_height);sh.fill.solid();sh.fill.fore_color.rgb=rgb(NAVY);sh.line.fill.background()
 txt(s,.7,.45,11,.3,section+' / MRM DIGITAL COMMUNITY PLATFORM',10,CYAN,True)
 s.shapes.add_picture(io.BytesIO(ASSETS[1]),Inches(13.5),Inches(.37),height=Inches(.38));s.shapes.add_picture(io.BytesIO(ASSETS[2]),Inches(14.45),Inches(.39),height=Inches(.32))
 txt(s,.7,1.04,14.6,.7,title,28,WHITE,True);txt(s,.72,1.87,14.5,.6,sub,14,MUTED)
 arrow(s,.7,8.35,15.3,8.35,LINE)
 txt(s,.7,8.55,3.8,.23,'MRM / ZYLAX TECHNOLOGIES',9,MUTED,target=3)
 txt(s,4.5,8.53,8.9,.27,'PROPOSED PRODUCT DESIGN · OPTIONAL PARTICIPATION · SEPARATE ROLE APPROVALS',8,MUTED,align=PP_ALIGN.CENTER)
 txt(s,14.1,8.5,1.2,.3,f'{len(P.slides):02d} / 24',11,CYAN,True,PP_ALIGN.RIGHT)
 return s

def card(s,x,y,w,h,label,body,c=CYAN,target=None):
 sh=box(s,x,y,w,h)
 compact=h<1.7
 txt(s,x+.18,y+.12,w-.36,.25 if compact else .34,label,10.5 if compact else 12,c,True,target=target)
 txt(s,x+.18,y+(.44 if compact else .67),w-.36,h-(.53 if compact else .8),body,12 if compact else 14,MUTED)
 if target:LINKS.append((sh,target))

def step(s,x,y,w,h,num,title,body,c=CYAN,target=None):
 sh=box(s,x,y,w,h,line=c);txt(s,x+.18,y+.15,.5,.3,str(num).zfill(2),12,c,True)
 compact=h<2.4
 txt(s,x+.18,y+(.48 if compact else .62),w-.36,.48 if compact else .65,title,14 if compact else 16,WHITE,True,target=target)
 txt(s,x+.18,y+(1.05 if compact else 1.46),w-.36,h-(1.16 if compact else 1.59),body,12 if compact else 13,MUTED)
 if target:LINKS.append((sh,target))

def notes(s,data):s.notes_slide.notes_text_frame.text=json.dumps(data,indent=2,ensure_ascii=False)

s=base('Connect people’s reasons to use the platform','A person should understand why another feature helps them—and choose whether to cross into it.','CONNECTED PLATFORM STRATEGY')
txt(s,.85,2.78,13.9,1.2,'“I came here to teach. Why would I volunteer?”',32,WHITE,True)
card(s,.7,4.4,4.6,2.5,'START WITH A REAL PERSON','Their current role, goal and reason to consider a second service.')
card(s,5.7,4.4,4.6,2.5,'MAKE THE HANDOFF USEFUL','A relevant invitation, consent and a clear next action—not just a menu link.')
card(s,10.7,4.4,4.6,2.5,'SHOW VALUE ON BOTH SIDES','Explain what the person gains and what the recipient or service gains.',GREEN)
txt(s,.85,7.37,14.1,.6,'15 proposed bridges · 2 complete journeys · role boundaries · implementation priorities',17,GOLD,True)
notes(s,{'purpose':'Supplement the existing feature explanations with actor motivation and reciprocal value.','status':'All new cross-feature invitations and rules are proposals requiring product validation.'})

s=base('A Quran teacher chooses a second way to serve','Learning centre → voluntary support participation → a subject-relevant offer a seeker can assess.')
for i,(t,b) in enumerate([('Teacher’s reason','“I want to use my Quran expertise to support someone beyond a scheduled class.”'),('Bridge experience','Invite the teacher to offer Quran support; reuse consented expertise and complete volunteer checks.'),('Seeker’s reason','“This approved volunteer is also an approved Quran teacher. I can see why the match is relevant.”')]):
 x=.7+i*5;step(s,x,2.8,4.6,3.05,i+1,t,b,target=5)
 if i<2:arrow(s,x+4.65,4.3,x+4.95,4.3)
card(s,.7,6.18,7.1,1.63,'DOES EVERY TEACHER NEED VOLUNTEERING?','No. It becomes relevant when they choose to offer support.',GREEN)
card(s,8.2,6.18,7.1,1.63,'WHAT DOES THE TEACHER BADGE MEAN?','Relevant approved expertise—not guaranteed safety or automatic case access.',GOLD)

s=base('How the user journeys connect','Read each row as a reason to move between services. Select an outlined card to explore its bridge.','CONNECTED USER JOURNEYS')
rows=[('TEACHER', [('Teach a topic','Already using learning',5),('Offer relevant help','Opt in + volunteer approval',5),('Support a seeker','Relevant credentials visible',7)]),('VOLUNTEER',[('Identify a skill gap','A real duty needs preparation',6),('Take useful training','Assess readiness separately',6),('Return to service','Approved for suitable work',17)]),('SEEKER',[('Ask for guidance','Learning need is expressed',7),('Choose a suitable class','Consent + language + level',7),('Request continuity','Optional mentoring',9)]),('LEARNER',[('Discover a programme','From an approved topic',12),('Build confidence','Learn at their chosen pace',16),('Choose to contribute','Apply separately to volunteer',16)]),('MRM TEAM',[('See aggregate demand','Protect individual identities',13),('Commission learning','Match teachers to a real gap',13),('Refer willing people','Measure usefulness, not pressure',7)])]
for r,(role,nodes) in enumerate(rows):
 y=2.68+r*1.02;txt(s,.75,y+.25,1.7,.4,role,11,GOLD,True)
 for j,(t,b,p) in enumerate(nodes):
  x=2.6+j*4.32;box(s,x,y,3.95,.85,line=CYAN);txt(s,x+.14,y+.1,3.67,.27,t,12,WHITE,True,target=p);txt(s,x+.14,y+.48,3.67,.23,b,9.5,MUTED)
  if j<2:arrow(s,x+3.99,y+.43,x+4.27,y+.43)
notes(s,{'reading':'Each row shows actor motivation and an actionable handoff. Nodes are not unrestricted data access.'})

s=base('A connection is useful only when the person needs it','Feature completeness does not require forcing every person through every service.')
card(s,.7,2.85,7.1,2.15,'YES—WHEN THERE IS A REAL REASON','A teacher wants to contribute; a volunteer needs preparation; a seeker asks for learning help.',GREEN)
card(s,8.2,2.85,7.1,2.15,'NO—WHEN THE SECOND SERVICE ADDS NO VALUE','A teacher only wants to teach; a reader only wants Quran; a seeker needs immediate practical help.',GOLD)
for i,(t,b) in enumerate([('Relevant invitation','Why this action, for this person, at this point?'),('Voluntary transition','Can they decline without losing their original service?'),('Useful context','What chosen information saves them from starting again?'),('Bounded authority','Which new checks apply before the next role or record opens?')]):
 x=.7+(i%2)*7.5;y=5.37+(i//2)*1.2;card(s,x,y,7.1,1.05,t,b)

for b in BRIDGES:
 s=base(b['title'],b['pair']+'  |  '+b['person'],'PROPOSED USER JOURNEY '+str(b['page']-4).zfill(2))
 txt(s,.85,2.55,14.1,.5,b['need'],17,WHITE,True)
 card(s,.7,3.18,7.1,1.04,'USE WHEN',b['yes'],GREEN)
 card(s,8.2,3.18,7.1,1.04,'NOT REQUIRED WHEN',b['no'],GOLD)
 for i,(t,d) in enumerate(b['steps']):
  x=.7+i*3.74;step(s,x,4.48,3.36,1.85,i+1,t,d)
  if i<3:arrow(s,x+3.4,5.4,x+3.69,5.4)
 # compact operational details under the behavioural journey
 for i,(label,body,c) in enumerate([('CARRY FORWARD',b['carry'],CYAN),('KEEP SEPARATE',b['separate'],GOLD),('VALUE ON BOTH SIDES',b['benefit'],GREEN)]):
  card(s,.7+i*5,6.56,4.6,1.12,label,body,c)
 txt(s,.85,7.87,14.2,.33,'SUCCESS: '+b['signal'],10,MUTED)
 notes(s,b)

s=base('Teacher journey: teaching becomes an optional way to serve','One account can support two approved roles. The person chooses whether to use both.','COMPLETE USER JOURNEY')
journey=[('Arrive to teach','Apply through the learning centre.'),('Approved expertise','MRM approves subject and teacher status.'),('Optional invitation','“Would you like to offer support in Quran?”'),('Apply as volunteer','Agree availability and complete separate checks.'),('Relevant match','A permitted offer shows approved expertise.'),('Accepted support','Seeker accepts; limited case access opens.'),('Continue or pause','Teacher controls capacity; coordinator reviews outcomes.')]
for i,(t,b) in enumerate(journey):
 if i<4:x=.7+i*3.74;y=2.75
 else:x=.7+(i-4)*3.74;y=5.32
 step(s,x,y,3.36,2.1,i+1,t,b)
 if i<3 or 4<=i<6:arrow(s,x+3.4,y+1.05,x+3.69,y+1.05)
arrow(s,13.6,4.9,13.6,5.08);arrow(s,13.6,5.08,2.38,5.08);arrow(s,2.38,5.08,2.38,5.3)
txt(s,12.1,5.75,3.1,1.5,'No volunteering?\nTeaching continues normally.',19,GOLD,True)
notes(s,{'teacher_subscription':'Teacher publishing entitlement remains separate; it does not grant volunteer approval, case access or guaranteed matching.'})

s=base('Seeker journey: support can open a chosen learning path','Learning complements assistance when useful. It is never a condition for receiving help.','COMPLETE USER JOURNEY')
journey=[('Ask for help','A seeker requests beginner Quran guidance.'),('Appropriate helper','MRM matches an approved volunteer; teacher expertise may be relevant.'),('Choose learning','The seeker accepts an optional referral.'),('Suitable programme','Topic, level, language and timing match the chosen goal.'),('Learn and follow up','Attend; request mentoring only if useful.'),('Continue independently','Return to Quran tools and the selected practice.'),('Contribute if desired','A later volunteer application remains separate and optional.')]
for i,(t,b) in enumerate(journey):
 x=.7+(i%4)*3.74;y=2.75+(i//4)*2.57;step(s,x,y,3.36,2.1,i+1,t,b)
 if i<3 or 4<=i<6:arrow(s,x+3.4,y+1.05,x+3.69,y+1.05)
arrow(s,13.6,4.9,13.6,5.08);arrow(s,13.6,5.08,2.38,5.08);arrow(s,2.38,5.08,2.38,5.3)
txt(s,12.1,5.75,3.1,1.5,'Declines learning?\nSupport stays available under normal eligibility.',17,GOLD,True)

s=base('Build the strongest bridges first','Sequence by user value and dependence on agreed policy—not by drawing the largest possible network.','PRODUCT PRIORITIES')
for col,(label,c,why,keys) in enumerate([
 ('FIRST',GREEN,'Clear service handoffs and role boundaries',['teacher_support','volunteer_training','seeker_class','mentoring_support','specialist_referral','website_journey']),
 ('NEXT',CYAN,'Relevant discovery and development paths',['quran_class','learning_mentor','event_volunteer','feed_learning','capacity_learning','learner_volunteer']),
 ('LATER',GOLD,'Needs stronger operational evidence or consent design',['consented_story','prayer_event','recognition_growth'])]):
 x=.7+col*5;box(s,x,2.77,4.6,5.1);txt(s,x+.2,2.97,4.2,.35,label,15,c,True);txt(s,x+.2,3.48,4.2,.8,why,14,MUTED)
 for i,key in enumerate(keys):
  b=BY_KEY[key];txt(s,x+.2,4.45+i*.5,4.2,.4,f'{b["page"]:02d}  '+b['pair'],10.5,WHITE,target=b['page'])
notes(s,{'priority':'Editorial recommendation, not an agreed delivery commitment. Validate with teachers, volunteers, seekers and MRM operations.'})

s=base('Reuse helpful context; grant each role separately','One person can be a learner, teacher and volunteer. Each relationship still has its own approval and permissions.','SHARED PRODUCT RULES')
for i,(t,b,c) in enumerate([
 ('ONE PERSON','Reuse an account and consented profile details. Do not create a new identity for every feature.',CYAN),
 ('SEPARATE ROLES','Teacher approval, publishing entitlement and volunteer eligibility are independent states.',GOLD),
 ('EXPLICIT RELATIONSHIPS','An enrolment, mentor acceptance, duty confirmation or case assignment gives a bounded relationship.',GREEN),
 ('VISIBLE RELEVANCE','Show approved subject expertise and status on a support offer—not private learning or worship records.',CYAN),
 ('CONTROLLED CONTEXT','Carry selected goals, public programme IDs or referral status. Record who consented to what.',GOLD),
 ('SAFE EXIT AND RETURN','Decline, pause or leave the second journey without losing the original service; handle revoked roles.',GREEN)]):
 x=.7+(i%2)*7.5;y=2.8+(i//2)*1.63;card(s,x,y,7.1,1.4,t,b,c)
notes(s,{'suggested_records':['person','role approval','subject credential','volunteer eligibility','programme enrolment','mentor relationship','case assignment','consented referral'],'server_rule':'Recheck current status at display, acceptance and record access. A badge is descriptive, not an authorisation token.'})

s=base('Validate that each bridge earns its place','Agree the product behaviour, then test it with the people who would actually cross between services.','DECISIONS AND ACCEPTANCE')
card(s,.7,2.8,7.1,2.12,'PRODUCT DECISIONS','Confirm visible teacher credentials, volunteer checks, invitation timing, referral ownership and consent wording.',CYAN)
card(s,8.2,2.8,7.1,2.12,'USER VALUE TEST','Can a teacher, volunteer or seeker explain why the suggested second service helps—and decline it easily?',GREEN)
card(s,.7,5.22,7.1,2.12,'MEANINGFUL FLOW TESTS','Declining keeps teaching available. Payment grants no case access. Revoked credentials stop appearing as current. Expired links recover honestly.',GOLD)
card(s,8.2,5.22,7.1,2.12,'RECIPROCAL OUTCOME','Measure appropriate accepted handoffs and user-reported usefulness. Do not optimise pressure to enrol, volunteer or disclose.',GREEN)
txt(s,.85,7.68,14.1,.4,'A connected platform helps a person take the right next step for their goal.',19,WHITE,True)
assert len(P.slides)==24
for sh,page in LINKS:sh.click_action.target_slide=P.slides[page-1]
P.save(ROOT/'MRM_User_Value_Bridges_24_Pages.pptx')
(ROOT/'Bridge_Model.json').write_text(json.dumps(BRIDGES,indent=2,ensure_ascii=False))
(ROOT/'Contents.json').write_text(json.dumps(TITLES,indent=2))
(ROOT/'Text_Geometry.json').write_text(json.dumps(TEXTCHECK,indent=2))
print('Built',len(P.slides),'pages and',len(BRIDGES),'proposed bridges')
