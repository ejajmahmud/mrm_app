"""Prepend the 24-page user-journey section, preserving editable shapes and links."""
from pathlib import Path
from copy import deepcopy
import io,json
from pptx import Presentation
from pptx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
base=Presentation(ROOT.parent/'MRM_Strategic_Platform_72_Pages.pptx')
intro=Presentation(ROOT/'MRM_User_Value_Bridges_24_Pages.pptx')
assert base.slide_width==intro.slide_width and base.slide_height==intro.slide_height
old=list(base.slides);new=[base.slides.add_slide(base.slide_layouts[6]) for _ in intro.slides]
slide_map={s.part.partname:t for s,t in zip(intro.slides,new)}
for source,dest in zip(intro.slides,new):
 for shape in source.shapes:
  element=deepcopy(shape.element)
  for item in element.iter():
   for key in [qn('r:id'),qn('r:embed'),qn('r:link')]:
    rid=item.get(key)
    if not rid:continue
    rel=source.part.rels[rid]
    if rel.reltype.endswith('/image'):
     _,newrid=dest.part.get_or_add_image_part(io.BytesIO(rel.target_part.blob))
    elif rel.reltype.endswith('/slide'):
     newrid=dest.part.relate_to(slide_map[rel.target_part.partname].part,rel.reltype)
    elif rel.is_external:newrid=dest.part.relate_to(rel.target_ref,rel.reltype,is_external=True)
    else:raise ValueError(rel.reltype)
    item.set(key,newrid)
  dest.shapes._spTree.insert_element_before(element,'p:extLst')
 dest.notes_slide.notes_text_frame.text=source.notes_slide.notes_text_frame.text
# IDs/parts remain stable, preserving all existing original slide destinations.
ids=list(base.slides._sldIdLst)
for el in ids:base.slides._sldIdLst.remove(el)
for el in ids[72:]+ids[:72]:base.slides._sldIdLst.append(el)
assert len(base.slides)==96
base.save(ROOT/'MRM_Complete_Connected_Platform_96_Pages.pptx')
# Roundtrip: all links target a slide in the combined presentation.
check=Presentation(ROOT/'MRM_Complete_Connected_Platform_96_Pages.pptx');valid={s.slide_id for s in check.slides}
links=0
for s in check.slides:
 for sh in s.shapes:
  t=sh.click_action.target_slide
  if t is not None:assert t.slide_id in valid;links+=1
print('Combined 96 editable slides;',links,'valid internal links')
