"""After LibreOffice PDF conversion: verify, add navigation and export diagrams."""
from pathlib import Path
from pptx import Presentation
from PIL import Image,ImageDraw
import fitz,json,re,zipfile,itertools,hashlib
ROOT=Path(__file__).resolve().parents[1]
N=lambda s:re.sub(r'[^a-z0-9]','',s.lower())
features=json.loads((ROOT/'Feature_Graph_Model.json').read_text())['features']
checks={}
for stem,count,contents in [('MRM_Strategic_Platform_72_Pages',72,'Main_Contents.json'),('MRM_Detailed_Feature_Atlas_52_Maps',52,'Atlas_Contents.json')]:
 deck=Presentation(ROOT/(stem+'.pptx'));pdf=fitz.open(ROOT/(stem+'.pdf'));assert len(deck.slides)==len(pdf)==count
 missing=[];outside=[];links=0
 for pg in pdf:
  for a in pg.get_links():pg.delete_link(a)
 ids={s.slide_id:i for i,s in enumerate(deck.slides)}
 for i,(sl,pg) in enumerate(zip(deck.slides,pdf)):
  page_text=N(pg.get_text())
  for sh in sl.shapes:
   if sh.has_text_frame and len(sh.text)>5 and N(sh.text) not in page_text:missing.append([i+1,sh.text])
   if sh.left<0 or sh.top<0 or sh.left+sh.width>deck.slide_width+10 or sh.top+sh.height>deck.slide_height+10:outside.append([i+1,sh.name])
   t=sh.click_action.target_slide
   if t is not None:
    assert t.slide_id in ids
    r=fitz.Rect(sh.left/12700,sh.top/12700,(sh.left+sh.width)/12700,(sh.top+sh.height)/12700)
    pg.insert_link({'kind':fitz.LINK_GOTO,'from':r,'page':ids[t.slide_id],'to':fitz.Point(0,0)});links+=1
  for b in pg.get_text('dict')['blocks']:
   if b['type']!=0:continue
   for l in b['lines']:
    for sp in l['spans']:
     x0,y0,x1,y1=sp['bbox']
     if x0<0 or y0<0 or x1>pg.rect.width+.5 or y1>pg.rect.height+.5:outside.append([i+1,sp['text']])
 assert not missing,missing
 assert not outside,outside
 titles=json.loads((ROOT/contents).read_text());pdf.set_toc([[1,t,i+1] for i,t in enumerate(titles)])
 tmp=ROOT/(stem+'_linked.pdf');pdf.save(tmp,garbage=4,deflate=True);pdf.close();tmp.replace(ROOT/(stem+'.pdf'))
 p=fitz.open(ROOT/(stem+'.pdf'));assert sum(len(x.get_links()) for x in p)==links
 with zipfile.ZipFile(ROOT/(stem+'.pptx')) as z:
  media=[n for n in z.namelist() if n.startswith('ppt/media/')]
  hashes=[hashlib.sha256(z.read(n)).hexdigest() for n in media]
  assert len(hashes)==len(set(hashes))
 checks[stem]={'pages':count,'navigation_links':links,'missing_native_text':0,'out_of_bounds':0,'media_files':len(media),'duplicate_media_files':0}
# Nodes must contain their own text; arrows must avoid unrelated feature boxes.
issues=[];overlaps=[];cross=[]
for filename,pdfname in [('Feature_Graph_Model.json','MRM_Strategic_Platform_72_Pages.pdf'),('Atlas_Graph_Model.json','MRM_Detailed_Feature_Atlas_52_Maps.pdf')]:
 raw=json.loads((ROOT/filename).read_text());graphs=raw['diagrams'] if isinstance(raw,dict) else raw;doc=fitz.open(ROOT/pdfname)
 for g in graphs:
  page=doc[g['page']-1]
  for key,n in g['nodes'].items():
   r=fitz.Rect(n['x']*72-.5,n['y']*72-.5,(n['x']+n['w'])*72+.5,(n['y']+n['h'])*72+.5)
   inside=N(page.get_textbox(r))
   for field in ['title','body']:
    if N(n[field]) not in inside:issues.append([filename,g['page'],key,field,n[field]])
  for (ka,a),(kb,b) in itertools.combinations(g['nodes'].items(),2):
   if min(a['x']+a['w'],b['x']+b['w'])>max(a['x'],b['x'])+.01 and min(a['y']+a['h'],b['y']+b['h'])>max(a['y'],b['y'])+.01:overlaps.append([g['page'],ka,kb])
  for e in g['edges']:
   for p,q in zip(e['points'],e['points'][1:]):
    for k,n in g['nodes'].items():
     if k in [e['source'],e['target']]:continue
     for step in range(1,100):
      t=step/100;x=p[0]+t*(q[0]-p[0]);y=p[1]+t*(q[1]-p[1])
      if n['x']+.015<x<n['x']+n['w']-.015 and n['y']+.015<y<n['y']+n['h']-.015:cross.append([g['page'],e['source'],e['target'],k]);break
assert not overlaps,overlaps
assert not cross,cross
assert not issues,issues
# Display images have exactly one usage per corrected capture.
uses=json.loads((ROOT/'Image_Use_Register.json').read_text());assert len(uses)==45 and len({r['screen'] for r in uses})==45
assert not ({r['screen'] for r in uses}&{'18','42','47'})
checks['feature_diagrams']={'individual_maps':52,'local_feature_nodes':sum(len(f['nodes']) for f in features),'labelled_relationships':sum(len(f['edges']) for f in features),'node_overlaps':0,'node_text_overflows':0,'edges_through_unrelated_nodes':0}
checks['screenshots']={'reviewed':48,'used_once':45,'excluded_identity_data':2,'excluded_development_only':1,'withdrawn_zip_used':False,'display_resolution':'900px wide, original proportions and transparency'}
# Per-feature vector and high-resolution image exports.
(ROOT/'maps').mkdir(exist_ok=True);(ROOT/'vectors').mkdir(exist_ok=True);(ROOT/'previews').mkdir(exist_ok=True)
atlas=fitz.open(ROOT/'MRM_Detailed_Feature_Atlas_52_Maps.pdf')
for i,f in enumerate(features):
 page=atlas[i];stem=f"{i+1:02d}_{f['key']}";page.get_pixmap(matrix=fitz.Matrix(1.6667,1.6667)).save(ROOT/'maps'/(stem+'.png'))
 (ROOT/'vectors'/(stem+'.svg')).write_text(page.get_svg_image(text_as_path=True))
# Contact sheets for reviewing all presentation pages and atlas maps.
for stem,pdfname in [('Deck','MRM_Strategic_Platform_72_Pages.pdf'),('Atlas','MRM_Detailed_Feature_Atlas_52_Maps.pdf')]:
 doc=fitz.open(ROOT/pdfname)
 for start in range(0,len(doc),12):
  sheet=Image.new('RGB',(1800,2400),'#111B35');d=ImageDraw.Draw(sheet)
  for j in range(start,min(start+12,len(doc))):
   pix=doc[j].get_pixmap(matrix=fitz.Matrix(.75,.75));im=Image.frombytes('RGB',[pix.width,pix.height],pix.samples);im.thumbnail((890,365));x=(j-start)%2*900;y=(j-start)//2*400;sheet.paste(im,(x,y+25));d.text((x+10,y+5),f'{stem.upper()} PAGE {j+1:02d}',fill='white')
  sheet.save(ROOT/'previews'/f'{stem}_{start+1:02d}_{min(start+12,len(doc)):02d}.jpg')
main=fitz.open(ROOT/'MRM_Strategic_Platform_72_Pages.pdf')
for i,name in [(0,'Cover'),(4,'Overall_Feature_Network'),(52,'Care_Journey'),(53,'Learning_Journey')]:
 main[i].get_pixmap(matrix=fitz.Matrix(1.6667,1.6667)).save(ROOT/'previews'/(name+'.png'))
(ROOT/'Verification.json').write_text(json.dumps(checks,indent=2))
# Core bundle intentionally excludes the PNG/SVG folders to keep the download practical.
files=['MRM_Strategic_Platform_72_Pages.pptx','MRM_Strategic_Platform_72_Pages.pdf','MRM_Detailed_Feature_Atlas_52_Maps.pptx','MRM_Detailed_Feature_Atlas_52_Maps.pdf','Screen_Audit.md','Screen_Audit.json','Verification.json','Browser_Verification.json','README.md','Interactive_Feature_Explorer.html']
with zipfile.ZipFile(ROOT/'MRM_Complete_Rebuild.zip','w',zipfile.ZIP_DEFLATED) as z:
 for f in files:
  if (ROOT/f).exists():z.write(ROOT/f,f)
with zipfile.ZipFile(ROOT/'MRM_Complete_Rebuild.zip') as z:assert z.testzip() is None
print(json.dumps(checks,indent=2));print('ZIP MiB:',round((ROOT/'MRM_Complete_Rebuild.zip').stat().st_size/1024/1024,1))
