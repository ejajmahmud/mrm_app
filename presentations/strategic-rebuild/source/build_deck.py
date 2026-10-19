"""Build a 72-page evidence-led MRM deck and a 52-map editable atlas.
Use only the corrected Google Drive folder, never the withdrawn screenshot ZIP.
Set MRM_SCREENSHOT_DIR to the extracted corrected image directory.
"""
from pathlib import Path
import os,io,json,math,re,hashlib
from PIL import ImageFont,Image
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from feature_content import FEATURES,CORE,EXTENSIONS,BY_KEY

ROOT=Path(__file__).resolve().parents[1]
SRC=Path(os.environ.get('MRM_SCREENSHOT_DIR','/workspace/scratch/mrm-rebuild/corrected-screens'))
OLD=ROOT.parent/'MRM_50_Page_Deck.pptx'
ref=Presentation(OLD);ASSETS=[ref.slides[0].shapes[n].image.blob for n in [0,4,5]]
NAVY='111B35';PANEL='1B2846';LINE='34435F';WHITE='F7F9FF';MUTED='BBC8DF';CYAN='62DAE8';PURPLE='C2A3FF';GREEN='8DD8B1';GOLD='E8C578';RED='EEA5AB';FONT='DejaVu Sans'
BOLD=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf',200)
REG=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',200)
P=None;ATLAS=False;LINKS=[];USES=[];TITLES=[];GRAPHS=[];NODES={}
ALLFILES={p.name[:2]:p for p in sorted(SRC.glob('*.png'))}
DISPLAY=Path('/workspace/scratch/mrm-rebuild/display-copies');DISPLAY.mkdir(parents=True,exist_ok=True)
for sid,original_path in ALLFILES.items():
 if sid in ['18','42','47']:continue
 dest=DISPLAY/original_path.name
 if dest.exists() and dest.stat().st_mtime>=original_path.stat().st_mtime:continue
 im=Image.open(original_path);height=round(im.height*900/im.width);im.resize((900,height),Image.Resampling.LANCZOS).save(dest,optimize=True)
assert len(ALLFILES)==48,'All 48 corrected screenshots are required for the audit.'
CAPTIONS={
'01':'Brand / loading screen','02':'Onboarding · care pillar','03':'Welcome and guest entry','04':'Registration · participation tracks','05':'Sign-in form','06':'Admin Ops Home','07':'Menu and utility directory','08':'Analytics + calculator concept','09':'Search: utilities and local places','10':'Infaq & Relief Campaigns','11':'Campaign detail · Overview','12':'Profile / role summary','13':'Settings, zone, lock and cache','14':'Halaqa & duty calendar','15':'Dawah & Impact Feed','16':'Caseworker conversation','17':'Sadaqah & Infaq Allocation','18':'Payment form · not reproduced','19':'In-app notification list','20':'Loading-state skeleton','21':'Waqt, fasting and prayer windows','22':'Private Salah / Qaza tracker','23':'Compass simulation + place list','24':'Al-Fatihah text / reciter controls','25':'Digital tasbih counter','26':'Dua categories and references','27':'Salah guide · posture 1 of 7','28':'Asma-ul-Husna names / meanings','29':'Ingredient reference + Fitrah concept','30':'Volunteer on-duty workspace','31':'Seeker workspace','32':'Scholar / Mentor roster','33':'Applicant verification progress','34':'Member workspace','35':'Onboarding · welfare / aid pillar','36':'Onboarding · volunteer pillar','37':'Campaign detail · Provisions','38':'Campaign detail · sample Impact','39':'Calendar almanac / prayer schedule','40':'Slot-confirmation modal','41':'Community-post composer','42':'Receipt concept · not reproduced','43':'Empty-state panel','44':'Network-error panel','45':'Role-picker modal','46':'Screen-directory modal','47':'Developer installation help · excluded','48':'Admin Home · dark-theme state'}

def rgb(c):return RGBColor.from_string(c)
def fit(txt,w,maxsize=16,bold=False):
 f=BOLD if bold else REG
 return min(maxsize,maxsize*w*72/(f.getlength(txt)/(200/maxsize))*0.94) if txt else maxsize

def box(s,x,y,w,h,fill=PANEL,line=None,round=True,alpha=None):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill)
 if line:sh.line.color.rgb=rgb(line);sh.line.width=Pt(.9)
 else:sh.line.fill.background()
 if round:
  try:sh.adjustments[0]=.10
  except:pass
 if alpha is not None:
  e=OxmlElement('a:alpha');e.set('val',str(alpha));sh.fill._xPr.solidFill.srgbClr.append(e)
 return sh

def txt(s,x,y,w,h,content,sz=16,col=WHITE,bold=False,align=PP_ALIGN.LEFT):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.clear();tf.word_wrap=True;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 for i,l in enumerate(content.split('\n')):
  pa=tf.paragraphs[0] if i==0 else tf.add_paragraph();pa.text=l;pa.font.name=FONT;pa.font.size=Pt(sz);pa.font.bold=bold;pa.font.color.rgb=rgb(col);pa.alignment=align;pa.space_after=Pt(2)
 return sh

def chip(s,x,y,w,label,col=CYAN,target=None):
 sh=box(s,x,y,w,.32,line=LINE);tx=txt(s,x+.1,y+.055,w-.2,.23,label,fit(label,w-.2,10,True),col,True)
 if target:LINKS.extend([(sh,target),(tx,target)])
 return sh

def base(title,sub,section='FEATURE CONNECTIONS'):
 global NODES
 NODES={};s=P.slides.add_slide(P.slide_layouts[6]);n=len(P.slides);TITLES.append(title)
 s.shapes.add_picture(io.BytesIO(ASSETS[0]),0,0,width=P.slide_width,height=P.slide_height)
 box(s,0,0,16,9,NAVY,round=False,alpha=90000)
 box(s,.7,.48,.045,.3,CYAN,round=False);txt(s,.9,.49,11,.27,section+'  /  MRM DIGITAL COMMUNITY PLATFORM',10,CYAN,True)
 s.shapes.add_picture(io.BytesIO(ASSETS[1]),Inches(13.5),Inches(.38),height=Inches(.38));s.shapes.add_picture(io.BytesIO(ASSETS[2]),Inches(14.45),Inches(.40),height=Inches(.32))
 txt(s,.7,1.02,14.6,.65,title,fit(title,14.6,30,True),WHITE,True)
 txt(s,.72,1.78,14.5,.6,sub,13,MUTED)
 box(s,.7,8.3,14.6,.012,LINE,round=False)
 tx=txt(s,.7,8.49,3.5,.25,'MRM  /  ZYLAX TECHNOLOGIES',9,MUTED);LINKS.append((tx,1 if ATLAS else 3))
 txt(s,4.1,8.49,8.0,.25,'PROPOSED CONNECTIONS  /  MVP UI DOES NOT VERIFY BACKEND BEHAVIOUR',8,MUTED,align=PP_ALIGN.CENTER)
 txt(s,14.1,8.43,1.2,.35,f'{n:02d} / {52 if ATLAS else 72}',12,CYAN,True,PP_ALIGN.RIGHT)
 return s

def band(s,label,x=.7,w=14.6,y=7.80,col=CYAN):
 box(s,x,y,w,.35);txt(s,x+.14,y+.065,w-.28,.24,label,fit(label,w-.28,10.3),col)

def textcard(s,x,y,w,h,title,body,col=CYAN,target=None):
 sh=box(s,x,y,w,h,line=LINE);box(s,x+.02,y+.16,.035,h-.32,col,round=False)
 tx=txt(s,x+.22,y+.20,w-.44,.55,title,fit(title,w-.44,18,True),WHITE,True)
 by=.66 if h<1.7 else .97
 txt(s,x+.22,y+by,w-.44,h-by-.10,body,13 if h<1.7 else 14,MUTED)
 if target:LINKS.extend([(sh,target),(tx,target)])

def node(s,key,x,y,w,h,title,body='',col=CYAN,target=None,focus=False):
 sh=box(s,x,y,w,h,fill='233B55' if focus else PANEL,line=col if focus else LINE)
 box(s,x+.02,y+.13,.035,h-.26,col,round=False)
 fs=fit(title,w-.36,16 if w>=3.4 else 14.5,True)
 tx=txt(s,x+.18,y+.16,w-.36,.45,title,fs,WHITE,True)
 txt(s,x+.18,y+.66,w-.36,h-.73,body,12.5 if w>=3.4 else 11,MUTED)
 if target:LINKS.extend([(sh,target),(tx,target)])
 NODES[key]={'x':x,'y':y,'w':w,'h':h,'title':title,'body':body,'target':target}

def an(k,side):
 a=NODES[k];x,y,w,h=(a[c] for c in ['x','y','w','h'])
 return {'L':(x,y+h/2),'R':(x+w,y+h/2),'T':(x+w/2,y),'B':(x+w/2,y+h)}[side]

def stroke(s,points,col=CYAN,dash=False,arrow=True):
 for j,(a,b) in enumerate(zip(points,points[1:])):
  sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(a[0]),Inches(a[1]),Inches(b[0]),Inches(b[1]));sh.line.color.rgb=rgb(col);sh.line.width=Pt(1.4)
  if dash:
   e=OxmlElement('a:prstDash');e.set('val','dash');sh.line._get_or_add_ln().append(e)
  if arrow and j==len(points)-2:
   e=OxmlElement('a:tailEnd');e.set('type','triangle');e.set('w','sm');e.set('len','sm');sh.line._get_or_add_ln().append(e)
  el=sh._element;el.getparent().remove(el);s.shapes._spTree.insert(4,el)

def connect(s,a,b,label,col=CYAN,dash=False,manual=None):
 aa=NODES[a];bb=NODES[b];ac=(aa['x']+aa['w']/2,aa['y']+aa['h']/2);bc=(bb['x']+bb['w']/2,bb['y']+bb['h']/2)
 if manual:pts=manual
 elif abs(ac[1]-bc[1])<.01:
  if ac[0]<bc[0]:pts=[an(a,'R'),an(b,'L')]
  else:pts=[an(a,'L'),an(b,'R')]
 elif abs(ac[0]-bc[0])<.01:
  pts=[an(a,'B' if ac[1]<bc[1] else 'T'),an(b,'T' if ac[1]<bc[1] else 'B')]
 elif ac[1]<bc[1]:
  p=an(a,'B');q=an(b,'T');mid=(p[1]+q[1])/2
  pts=[p,(p[0],mid),(q[0],mid),q]
 else:
  sa='R' if ac[0]<bc[0] else 'L';sb='L' if ac[0]<bc[0] else 'R';p=an(a,sa);q=an(b,sb);mid=(p[0]+q[0])/2
  pts=[p,(mid,p[1]),(mid,q[1]),q]
 stroke(s,pts,col,dash)
 # Place each action on the longest route segment, away from its nodes.
 seg=max(zip(pts,pts[1:]),key=lambda z:abs(z[0][0]-z[1][0])+abs(z[0][1]-z[1][1]));u,v=seg;mx=(u[0]+v[0])/2;my=(u[1]+v[1])/2
 if abs(u[1]-v[1])<.01:
  avail=abs(u[0]-v[0]);w=min(1.15,max(.66,avail-.08));ly=my-.34
 else:w=min(1.24,max(.72,(aa['x']-bb['x'] if aa['x']>bb['x'] else bb['x']-aa['x'])-.1));ly=my-.13
 lx=mx-w/2;box(s,lx,ly,w,.27,fill=NAVY,round=False);txt(s,lx+.03,ly+.04,w-.06,.22,label,fit(label,w-.06,8.5,True),col,True,PP_ALIGN.CENTER)
 return {'source':a,'target':b,'meaning':label,'points':pts,'dashed':dash}

def phone(s,id,x,y=2.65,h=4.8):
 assert id not in ['18','42','47'],'Sensitive / development-only images must not be embedded.'
 original_path=ALLFILES[id];p=DISPLAY/original_path.name;im=Image.open(p);w=h*im.width/im.height
 s.shapes.add_picture(str(p),Inches(x),Inches(y),height=Inches(h))
 txt(s,x-.06,y+h+.08,w+.12,.44,'MVP '+id+' · '+CAPTIONS[id],9.2,MUTED,align=PP_ALIGN.CENTER)
 USES.append({'page':len(P.slides),'screen':id,'file':p.name,'actual_content':CAPTIONS[id]})
 return w

def target(ref):
 f=BY_KEY.get(ref)
 return (f['atlas_page'] if ATLAS else f['page']) if f else None

def feature(f):
 has=bool(f['screen']) and not ATLAS;n=len(f['screen'] or [])
 s=base(f['title'],f['why'],'EXTENDED CONCEPT' if f['kind']=='extension' else 'FEATURE NETWORK')
 if has:
  if n==1:gx=.7+3.7;gw=10.9;ph=4.82;px=.88
  else:gx=5.66;gw=9.64;ph=4.65;px=.7
  for i,sc in enumerate(f['screen']):phone(s,sc,px+i*2.42,2.65,ph)
 else:gx=.7;gw=14.6
 chip(s,gx,2.45,gw,f['status'],GOLD if f['kind']=='extension' else CYAN)
 gap=.75 if gw<11 else .84;nw=(gw-2*gap)/3;xx=[gx+i*(nw+gap) for i in range(3)];yy=[3.08,5.14];nh=1.34
 for i,d in enumerate(f['nodes']):node(s,str(i),xx[i%3],yy[i//3],nw,nh,d['title'],d['body'],PURPLE if f['kind']=='extension' else CYAN,target(d['ref']),i==1)
 es=[]
 for a,b,label in f['edges']:es.append(connect(s,str(a),str(b),label,GOLD if label in ['reports','escalates','expires','fails'] else (PURPLE if f['kind']=='extension' else CYAN)))
 box(s,gx,6.98,gw,.82,fill='253651',line=LINE)
 txt(s,gx+.18,7.08,gw-.36,.21,'CRITICAL CONNECTION TO GET RIGHT',9,GOLD,True)
 critical=' → '.join(f['nodes'][i]['title'] for i in [3,1,2])
 wfs=16 if REG.getlength(critical)/200*16 <= (gw-.36)*72 else 14
 txt(s,gx+.18,7.27,gw-.36,.51,critical,wfs,WHITE)
 band(s,'Measure after implementation: '+f['measure'],gx,gw,7.89,MUTED)
 s.notes_slide.notes_text_frame.text=json.dumps({'feature':f['key'],'owner':f['owner'],'scope':f['status'],'why':f['why'],'measure':f['measure'],'decision':f['decision'],'screen_evidence':[CAPTIONS[x] for x in f['screen'] or []],'connections':[(f['nodes'][a]['title'],label,f['nodes'][b]['title']) for a,b,label in f['edges']]},indent=2)+'\nScreen UI is evidence of a visual state, not tested persistence, permissions, live APIs, religious approval, tax certification, payment or production impact. Only corrected Drive images are used.'
 GRAPHS.append({'feature':f['key'],'page':len(P.slides),'nodes':dict(NODES),'edges':es})
 return s


def init(atlas=False):
 global P,ATLAS,LINKS,USES,TITLES,GRAPHS
 P=Presentation();P.slide_width=Inches(16);P.slide_height=Inches(9);ATLAS=atlas;LINKS=[];USES=[];TITLES=[];GRAPHS=[]

def link_all():
 for sh,t in LINKS:
  assert 1<=t<=len(P.slides),(t,len(P.slides))
  sh.click_action.target_slide=P.slides[t-1]

# ---------- Main presentation: executive story, feature atlas, cross-feature systems, delivery ----------
init()
s=base('MRM: faith, learning and care, connected','A digital platform for all Muslims, with dedicated support for reverts and people who need assistance.','PLATFORM VISION')
txt(s,.75,2.65,10,.38,'Bismillahirrahmanirrahim',16,GOLD)
txt(s,.7,3.28,11,1.5,'Every feature should make\nthe next experience better.',35,WHITE,True)
txt(s,.75,5.18,10.8,.98,'Daily practice opens the door. Trusted learning builds participation.\nConnected service turns a need into accountable care.',20,MUTED)
phone(s,'48',12.55,2.65,4.75)
for i,(a,b,c) in enumerate([('FAITH','Useful every day',GOLD),('LEARNING','Approved knowledge',PURPLE),('CARE','Responsible follow-up',GREEN),('OPERATIONS','A service MRM can run',CYAN)]):
 x=.7+(i%2)*5.55;y=6.55+(i//2)*.59
 box(s,x,y,5.22,.46,line=c);txt(s,x+.15,y+.115,4.92,.28,a+'  /  '+b,12,c,True)
band(s,'A connected specification, supported by audited MVP visuals and explicit feature relationships.',w=11.0)
# 02
s=base('The value comes from the connections','The proposal creates an operating service across mobile apps, a public website, administration and shared backend.','WHY THIS PLATFORM')
for i,(a,b,c) in enumerate([
('Useful every day','Prayer, Quran and daily guidance give Muslims reasons to return beyond a one-off request.',GOLD),
('Trusted participation','Reviewed teachers, verified volunteers and controlled publication make the next action clearer.',PURPLE),
('Accountable care','A request leads to triage, a confirmed assignment, follow-up and an auditable outcome.',GREEN),
('Sustainable operations','Teacher entitlements and sponsor placements support revenue without gating core assistance.',CYAN)]):textcard(s,.7+(i%2)*7.45,2.7+(i//2)*2.38,7.15,2.13,a,b,c)
band(s,'Success should be measured through real participation, reliable service completion and MRM’s operational capacity.')
# 03
s=base('A presentation you can explore by outcome','Start with the overall network, then follow a feature into its detailed dependencies, users and value.','NAVIGATION')
for i,(a,b,c,t) in enumerate([
('01  Overall platform','Page 05\nThe principal cross-feature paths.',CYAN,5),('02  Individual feature maps','Pages 06–52\n47 primary capability networks and their importance.',GREEN,6),('03  End-to-end journeys','Pages 53–56\nCare, learning, daily use and event delivery.',PURPLE,53),('04  Extended concepts','Pages 57–61\nFive financial / impact concepts, clearly qualified.',GOLD,57),('05  Delivery & acceptance','Pages 62–65\nSequencing, ownership, checks and decisions.',CYAN,62),('06  Evidence & connection index','Pages 66–72\nSource review, traceability and next steps.',PURPLE,66)]):textcard(s,.7+(i%3)*4.94,2.65+(i//3)*2.36,4.69,2.14,a,b,c,t)
band(s,'Feature nodes and index entries are clickable in PowerPoint and PDF. Every footer returns here.')
# 04
s=base('All Muslims can benefit; specialised support stays focused','Different journeys share an identity foundation while keeping permissions and personal information separate.','AUDIENCE & VALUE')
for i,(a,b,c,t) in enumerate([
('Everyday Muslim users','Daily tools, approved resources and programme discovery.',GOLD,'prayer'),('Members / support seekers','Safe requests, personal updates and relevant guidance.',GREEN,'requests'),('Volunteers / applicants','A defined approval route and suitable service opportunities.',GREEN,'vetting'),('Teachers / mentors','Approved teaching and explicit publishing entitlements.',PURPLE,'teacher'),('MRM staff','Duty-based queues, decisions, oversight and controlled reports.',CYAN,'admin'),('Partners / sponsors','Approved referral coordination and accountable recognition.',GOLD,'sponsors')]):textcard(s,.7+(i%3)*4.94,2.65+(i//3)*2.36,4.69,2.14,a,b,c,target(t))
band(s,'A role is not a paywall. Paying, earning rank or choosing a demo persona does not create staff privileges.')
# 05: true feature-to-feature network with cross-lane events, no MRM hub.
s=base('The overall network: features hand off real work','Identity, publishing and participation converge through permission checks, accountable records and relevant updates.','OVERALL FEATURE NETWORK')
xs=[.7,4.55,8.4,12.25];ys=[2.72,4.48,6.24];w=3.05;h=1.14
items=[('a',0,0,'Identity & verified roles','Account, approval and visibility','roles'),('b',1,0,'Request & matching','Triage and eligible volunteer offer','matching'),('c',2,0,'Confirmed assignment','Permitted contact and progress','assignments'),('d',3,0,'Case / service history','Outcome, closure and follow-up','activity'),('e',0,1,'Approved teacher','Credentials and professional role','teacher'),('f',1,1,'Verified plan rights','Teacher-only publishing privileges','billing'),('g',2,1,'Moderated publishing','Approved posts and programmes','posts'),('h',3,1,'Classes & calendar','Registration, sessions and updates','schedule'),('i',0,2,'Daily tools / website','Relevant public entry points','home'),('j',1,2,'Search & discovery','Find approved content / services','search'),('k',2,2,'Join, apply or request','Take the relevant permitted action','requests'),('l',3,2,'Safe alerts & activity','Open the authorised next step','alerts')]
for key,c,r,a,b,refk in items:node(s,key,xs[c],ys[r],w,h,a,b,[GREEN,PURPLE,GOLD][r],target(refk))
for a,b,t in [('a','b','gates'),('b','c','confirms'),('c','d','records'),('e','f','qualifies'),('f','g','permits'),('g','h','publishes'),('i','j','opens'),('j','k','finds'),('k','l','triggers'),('a','e','binds'),('h','l','reminds')]:connect(s,a,b,t)
connect(s,'d','l','updates',CYAN,manual=[an('d','R'),(15.57,3.29),(15.57,6.81),an('l','R')])
band(s,'Cross-cutting checks: consent • role / record permissions • content approval • audit • recoverable state')
GRAPHS.append({'feature':'overall','page':5,'nodes':dict(NODES),'edges':[]})
# 06–52: 47 individual networks.
for f in CORE:
 assert len(P.slides)+1==f['page'];feature(f)
# 53: care system, with distinct onboarding image.
s=base('Care is a connected service, not a dispatch button','Verification, matching, contact, closure and recognition have distinct responsibilities and access boundaries.','END-TO-END CARE NETWORK')
phone(s,'36',.85,2.65,4.82)
for key,x,y,a,b,refk in [('a',4.45,3.05,'Private request','Triage selects the routine or restricted route','requests'),('b',8.35,3.05,'Eligible matching','Approved active volunteer + access rules','matching'),('c',12.25,3.05,'Confirmed assignment','Record who is responsible for the work','assignments'),('d',4.45,5.20,'Verification & interview','Review evidence before eligible service','vetting'),('e',8.35,5.20,'Permitted case contact','Progress, report concerns and follow up','chat'),('f',12.25,5.20,'Closure & recognition','Verify outcome before service credit','recognition')]:node(s,key,x,y,3.05,1.35,a,b,GREEN,target(refk))
for a,b,t in [('a','b','triages'),('b','c','offers'),('d','b','qualifies'),('c','f','confirms'),('c','e','permits'),('e','f','records')]:connect(s,a,b,t,GREEN)
band(s,'Restricted faith / family cases go to designated personnel, not the general volunteer broadcast route.',4.45,10.85)
# 54 learning system
s=base('Learning connects content, entitlement, time and participation','Teacher approval, payment verification and moderation are independent controls serving one learner journey.','END-TO-END LEARNING NETWORK')
for key,x,y,a,b,r in [('a',.7,2.85,'Teacher approval','MRM confirms credentials and role','teacher'),('b',5.85,2.85,'Plan entitlement','Verify designated publishing privileges','billing'),('c',11,2.85,'Content moderation','Religious / rights and publication review','posts'),('d',.7,5.13,'Discovery & guidance','Daily tools and public content open learning','search'),('e',5.85,5.13,'Programme registration','Schedule and access rules are explicit','schedule'),('f',11,5.13,'Session / recording','Permitted access, reminders and continuation','learning')]:node(s,key,x,y,4.3,1.4,a,b,PURPLE,target(r))
for a,b,t in [('a','b','qualifies'),('b','c','allows submission'),('c','f','publishes'),('d','e','registers'),('e','f','permits'),('c','e','lists')]:connect(s,a,b,t,PURPLE)
band(s,'No teacher payment grants private-case access. No proposed subscription paywall for ordinary members or volunteers.')
# 55 daily use pathway
s=base('Daily tools make the platform valuable between service events','A person can practise, learn and discover relevant participation without needing an assistance case.','EVERYDAY MUSLIM JOURNEY')
for key,x,y,a,b,r in [('a',.7,2.85,'Location / source settings','Approved local and content conventions','settings'),('b',5.85,2.85,'Daily Islamic utilities','Prayer, Quran, Qibla and remembrance','prayer'),('c',11,2.85,'Approved explanations','Guides, references and programmes','learning'),('d',.7,5.13,'Private personal choices','Consent, logs and supported cache','tracker'),('e',5.85,5.13,'Relevant discovery','Mosques, classes and support routes','search'),('f',11,5.13,'Voluntary participation','Join a programme or begin a request','schedule')]:node(s,key,x,y,4.3,1.4,a,b,GOLD,target(r))
for a,b,t in [('a','b','configures'),('b','c','explains'),('d','b','personalises'),('b','e','discovers'),('c','f','invites'),('e','f','continues')]:connect(s,a,b,t,GOLD)
band(s,'Personal worship logs stay private. Community participation is a choice, not a condition of using daily tools.')
# 56 events and oversight
s=base('One event connects state, notification and accountable follow-up','An alert is useful only when it points to the right record and preserves the current access boundary.','SHARED EVENT NETWORK')
for key,x,y,a,b,r in [('a',.7,2.85,'Application / case / class','A legitimate service state changes','activity'),('b',5.85,2.85,'Recipient & privacy checks','Role, record and notification preference','security'),('c',11,2.85,'Safe alert delivery','Minimal copy and relevant destination','alerts'),('d',.7,5.13,'Persistent event history','Record the actual transition and owner','admin'),('e',5.85,5.13,'Authorised next action','Recheck access when the user opens it','assignments'),('f',11,5.13,'Follow-up & reports','Review unattended work; aggregate outcomes','reports')]:node(s,key,x,y,4.3,1.4,a,b,CYAN,target(r))
for a,b,t in [('a','b','triggers'),('b','c','filters'),('a','d','records'),('c','f','prompts'),('c','e','opens'),('d','e','reconciles'),('e','f','updates')]:connect(s,a,b,t)
band(s,'Failed notification delivery does not erase the case; application state and staff follow-up remain authoritative.')
# 57–61 extended financial and impact concept maps, explicitly outside silent scope expansion.
for f in EXTENSIONS:
 assert len(P.slides)+1==f['page'];feature(f)
# 62 delivery
s=base('Delivery follows dependencies, not disconnected screen counts','The earlier proposal estimates 3–4 months for agreed scope, subject to required inputs and external readiness.','DELIVERY STRATEGY')
for i,(a,b,c) in enumerate([('Foundation','Source / licence audit, journeys, access matrix, content and acceptance specification.',CYAN),('Working service flows','Identity, verification, requests, teachers, learning and private records as complete vertical journeys.',GREEN),('Connected platform','Utilities, website, sponsors, notifications, reports and operational configuration.',PURPLE),('Acceptance & release','Meaningful permissions / device / billing checks, MRM UAT, training and client-owned submissions.',GOLD)]):textcard(s,.7+(i%2)*7.45,2.7+(i//2)*2.37,7.15,2.12,a,b,c)
band(s,'Extended collection / payment / receipt concepts require separate authority, technical assessment and written scope decisions.')
# 63 responsibilities
s=base('The software and the operating service need clear owners','MRM defines policies and owns its operational decisions; Zylax implements and verifies the agreed software.','OPERATING MODEL')
for i,(a,b,c) in enumerate([('MRM programme owners','Approval criteria, safeguarding policy, content rights, teacher terms and sponsor agreements.',GREEN),('MRM service operators','Case queues, interviews, follow-up, moderation, learning programmes and partner relationships.',GREEN),('Zylax delivery team','Template-consistent implementation, server permissions, persistence, testing, documentation and training.',CYAN),('Client-owned accounts','Hosting, domain, store verification, billing setup and third-party access with authorised collaboration.',PURPLE)]):textcard(s,.7+(i%2)*7.45,2.7+(i//2)*2.37,7.15,2.12,a,b,c)
band(s,'Store submission is a delivery milestone. Public approval and release dates remain controlled by the stores.')
# 64 acceptance
s=base('Acceptance should prove the connections actually work','Each check follows an end-to-end journey and attempts the important denied, failed and recovery paths.','ACCEPTANCE EVIDENCE')
checks=[('Identity → correct records','Recovery, role activation and direct API denial.'),('Applicant → eligible service','Evidence / interview / decision and immediate suspension.'),('Request → confirmed outcome','Triage, matching, assignment, reassignment and closure.'),('Teacher → permitted publication','Approval, billing state, allowance and independent moderation.'),('Programme → participant access','Registration, reminder, session and recording rights.'),('Confidential case → safe follow-up','Unrelated accounts, evidence files, exports and safe notifications.')]
for i,(a,b) in enumerate(checks):textcard(s,.7+(i%3)*4.94,2.65+(i//3)*2.36,4.69,2.14,a,b,[CYAN,GREEN,PURPLE][i%3])
band(s,'Acceptance follows proven journeys, permission boundaries and recoverable failures, with findings recorded and retested.')
# 65 dependency priorities
s=base('Why every capability belongs in the overall product','Different features create different kinds of value; their dependencies determine the implementation sequence.','IMPORTANCE & PRIORITIES')
for i,(a,b,c) in enumerate([('Trust foundation','Identity, permissions, approved sources, file access, reliable state and audit make other features safe to operate.',CYAN),('Service outcomes','Triage, verification, assignment and follow-up turn a need into an accountable outcome.',GREEN),('Everyday engagement','Prayer, Quran, guidance, search and programmes keep the platform useful for the wider Muslim community.',GOLD),('Sustainable growth','Teacher entitlements and sponsor delivery connect revenue to explicit privileges and accountable recognition.',PURPLE)]):textcard(s,.7+(i%2)*7.45,2.7+(i//2)*2.37,7.15,2.12,a,b,c)
band(s,'Sequence by dependency and operational readiness; keep approved modules visible rather than silently dropping difficult work.')
# 66 screenshot evidence scope
s=base('What the MVP images establish, and what they do not','All 48 corrected Drive captures were reviewed as UI states, not as independent proof of operational results.','SCREENSHOT EVIDENCE')
phone(s,'01',.85,2.7,4.76)
for i,(a,b,c) in enumerate([('Visible evidence','The images establish labels, screen structure, selected tabs, role states and the controls that are visible.',CYAN),('Not yet verified','Persistence, working APIs, live biometrics, encrypted chat, payments, tax status and operational impact need separate checks.',GOLD),('Religious / policy approval','Sources, local conventions, classifications, verification rules and safeguarding coverage require responsible approval.',GREEN)]):textcard(s,4.5,2.68+i*1.60,10.8,1.42,a,b,c)
band(s,'No statistics, audit badge, testimonial or “sync” label in a prototype is represented here as a proven result.',4.5,10.8)
# 67 source fidelity / development screens
s=base('Internal navigation is not a new technology commitment','The directory and menu are useful audit views; development helpers do not establish the delivered app framework.','SOURCE / ARCHITECTURE REVIEW')
phone(s,'07',.75,2.68,4.68);phone(s,'46',3.20,2.68,4.68)
textcard(s,6.0,2.7,9.3,2.12,'Preserve the chosen source strategy','The prior requirement is Zabi v4.3 / Flutter and template preservation. A menu saying Kotlin or a PWA helper is not approval to replace it.',CYAN)
textcard(s,6.0,5.06,9.3,2.12,'Audit before implementation claims','Inspect the actual source, licence, dependencies, routes, backend and platform targets before reconciling the prototype with release requirements.',GOLD)
band(s,'Capture 47 is developer installation help and is excluded from client feature imagery. No source architecture was validated here.')
# 68 audit register overview
s=base('Every corrected capture has one deliberate use or exclusion','Variants are identified by what changes on screen; no image is relabelled to imply a different completed feature.','IMAGE AUDIT & COVERAGE')
for i,(a,b,c) in enumerate([('48 images reviewed','Content labels, hashes, selected states and use locations are recorded in the accompanying screen register.',CYAN),('45 distinct visual uses','Distinct screens, roles, tabs, modals and theme states are used once each in this deck.',GREEN),('3 explicit exclusions','Payment form 18 and receipt 42 contain a full identity number. Image 47 is developer installation help.',GOLD),('Corrected explanations','Catalog / product / cart are relief-purpose views. Mentor roster is not subscription approval; a simulation is not a live sensor test.',PURPLE)]):textcard(s,.7+(i%2)*7.45,2.7+(i//2)*2.37,7.15,2.12,a,b,c)
band(s,'See Screen_Audit.json / Screen_Audit.md for the full file-to-page register and capture-specific findings.')
# 69 / 70 source-to-feature index
for part,chunk in enumerate([FEATURES[:26],FEATURES[26:]]):
 s=base('Feature connection index '+str(part+1)+' / 2','Choose an entry to open its individual network; each page explains the connection value and what to verify.','TRACEABLE FEATURE INDEX')
 for j,f in enumerate(chunk):
  col=j//13;row=j%13;x=.7+col*7.45;y=2.60+row*.385
  box(s,x,y,7.15,.35,fill=PANEL if row%2==0 else '17233D',round=False)
  tx=txt(s,x+.12,y+.065,6.62,.27,f"{f['page']:02d}   {f['key'].replace('_',' ').title()}",11,CYAN if f['kind']=='feature' else GOLD,True);LINKS.append((tx,f['page']))
 band(s,'The separate 52-map atlas uses the same connection model, with larger diagrams and its own feature navigation.')
# 71 decisions
s=base('Resolve the remaining decisions without inventing scope','A strong delivery specification separates agreed capabilities, operating configuration and genuinely new obligations.','DECISION RECORD')
for i,(a,b,c) in enumerate([('Source & platform','Confirm current Zabi source, licence, template reference and Android / iOS release architecture.',CYAN),('Service & teaching policies','Agree access exceptions, interview evidence, case escalation, mentor permissions, teacher plans and historical-content visibility.',GREEN),('External inputs','Confirm religious sources, partner sharing, mapping, streaming / recordings and applicable store billing.',PURPLE),('Extended financial concepts','Confirm collection authority, payment provider, reconciliation, disbursement and lawful receipt wording before a scope commitment.',GOLD)]):textcard(s,.7+(i%2)*7.45,2.7+(i//2)*2.37,7.15,2.12,a,b,c)
band(s,'Screen concepts inform discussion. They do not by themselves amend an approved SOW or establish authority access.')
# 72 close
s=base('A platform MRM can use, govern and grow','The goal is a coherent service for the wider Muslim community, with focused care and accountable delivery.','NEXT STEPS')
txt(s,.75,2.83,13.9,1.06,'Useful daily. Trusted in learning.\nResponsible in care.',33,WHITE,True)
for i,(a,b,c) in enumerate([('Confirm','Agree journeys, policy owners and acceptance criteria.',CYAN),('Implement','Build the connected flows with real permissions and persistence.',GREEN),('Demonstrate','Review evidence against meaningful end-to-end checks.',PURPLE),('Launch responsibly','Train MRM staff and submit through client-owned accounts.',GOLD)]):textcard(s,.7+i*3.70,4.68,3.50,2.45,a,b,c)
band(s,'Next step: nominate MRM’s decision makers and conduct the source, service-policy and implementation workshop.')
assert len(P.slides)==72,len(P.slides)
assert len(USES)==45,(len(USES),USES)
from collections import Counter
counts=Counter(x['screen'] for x in USES);assert all(n==1 for n in counts.values()),counts
assert set(counts)==set(ALLFILES)-{'18','42','47'},set(ALLFILES)-set(counts)
link_all();P.save(ROOT/'MRM_Strategic_Platform_72_Pages.pptx')
(ROOT/'Image_Use_Register.json').write_text(json.dumps(USES,indent=2))
(ROOT/'Main_Contents.json').write_text(json.dumps(TITLES,indent=2))
(ROOT/'Feature_Graph_Model.json').write_text(json.dumps({'features':FEATURES,'diagrams':GRAPHS},indent=2))
# A separate atlas enlarges each feature diagram and retains fully editable native shapes.
init(True)
for f in FEATURES:feature(f)
assert len(P.slides)==52;link_all();P.save(ROOT/'MRM_Detailed_Feature_Atlas_52_Maps.pptx')
(ROOT/'Atlas_Contents.json').write_text(json.dumps(TITLES,indent=2))
(ROOT/'Atlas_Graph_Model.json').write_text(json.dumps(GRAPHS,indent=2))
print('Built 72 presentation pages, 52 detailed feature maps, 45 unique corrected-screen uses.')
