"""Assemble a single client-facing proposal with a screenshot cover and global navigation."""
from pathlib import Path
from copy import deepcopy
import sys,io,json,os
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.oxml.ns import qn
from pptx.enum.shapes import MSO_SHAPE
ROOT=Path(__file__).resolve().parents[1]
HELPER=ROOT.parent/'strategic-rebuild/user-journeys/source/build_bridges.py'
sys.path.insert(0,str(HELPER.parent))
# Reuse the existing typography/layout functions without running its presentation build.
ns={'__file__':str(HELPER)}
exec(compile(HELPER.read_text().split("s=base('Connect people’s reasons")[0],str(HELPER),'exec'),ns)
P=ns['P'];txt=ns['txt'];box=ns['box'];card=ns['card'];arrow=ns['arrow'];base=ns['base'];notes=ns['notes'];ASSETS=ns['ASSETS']
WHITE=ns['WHITE'];MUTED=ns['MUTED'];CYAN=ns['CYAN'];GOLD=ns['GOLD'];GREEN=ns['GREEN'];PURPLE=ns['PURPLE'];NAVY=ns['NAVY']
DISPLAY=Path(os.environ.get('MRM_SCREENSHOT_DIR','/workspace/scratch/mrm-rebuild/display-copies'))
def image(s,sid,x,y,h):
 path=next(DISPLAY.glob(sid+'-*.png'));s.shapes.add_picture(str(path),Inches(x),Inches(y),height=Inches(h))

# A genuine proposal cover: actual app imagery, audience and presenter identity.
s=P.slides.add_slide(P.slide_layouts[6]);ns['TITLES'].append('MRM Digital Community Platform')
s.background.fill.solid();s.background.fill.fore_color.rgb=ns['rgb'](NAVY)
# Subtle architectural panels behind three distinct source screenshots.
box(s,8.25,.7,7.1,7.65,c='172541',line='263C59')
box(s,8.65,1.05,6.3,6.95,c='1B2C48',line='304B65')
s.shapes.add_picture(io.BytesIO(ASSETS[1]),Inches(.8),Inches(.7),height=Inches(.58))
s.shapes.add_picture(io.BytesIO(ASSETS[2]),Inches(5.32),Inches(.81),height=Inches(.35))
txt(s,.8,1.96,6.85,.4,'DIGITAL PLATFORM PROPOSAL',13,CYAN,True)
txt(s,.78,2.7,7.05,2.48,'MRM Digital\nCommunity\nPlatform',43,WHITE,True)
txt(s,.82,5.53,6.3,.84,'Faith. Learning. Care.\nConnected for the Muslim community.',21,WHITE)
txt(s,.82,6.77,6.3,.78,'For all Muslims, with dedicated support\nfor reverts and people seeking assistance.',14,MUTED)
image(s,'24',8.60,2.25,5.55)
image(s,'30',12.12,1.97,5.86)
image(s,'34',10.05,1.31,6.75)
txt(s,8.65,8.01,6.3,.25,'ACTUAL MVP APP SCREENS · PROPOSED PLATFORM EXPERIENCE',8.5,MUTED,align=ns['PP_ALIGN'].CENTER)
txt(s,.82,8.18,6.1,.42,'Prepared for MRM  |  Zylax Technologies  |  October 2026',10,MUTED)
notes(s,{'cover_sources':['24 Quran utility','30 volunteer workspace','34 member workspace'],'qualification':'Actual MVP screen states; imagery is illustrative, not evidence that production integrations or claims have been validated. Reuse on cover is intentional.'})

s=base('One platform for everyday value and community impact','Bring daily Muslim utilities, approved learning and responsible support into a coherent service.','EXECUTIVE PROPOSAL')
for i,(a,b,c) in enumerate([
 ('USEFUL EVERY DAY','Prayer, Quran, guidance and local discovery create everyday reasons to use the app.',GOLD),
 ('LEARNING WITH CONTINUITY','Approved teachers, programmes and mentoring help people take a meaningful next step.',PURPLE),
 ('CARE WITH ACCOUNTABILITY','Requests, volunteer approval, matching and follow-up connect people to appropriate help.',GREEN),
 ('OPERATIONS MRM CAN GOVERN','Administration, permissions, reporting and clear responsibilities support service delivery.',CYAN)]):
 card(s,.7+(i%2)*7.5,2.75+(i//2)*2.28,7.1,2.0,a,b,c)
txt(s,.87,7.63,14.1,.42,'Mobile app + public website + administration + shared service foundation',18,WHITE,True)

s=base('Designed for the whole Muslim community','Specialist revert support remains a focused service within a broader, inclusive platform.','WHO THE PLATFORM SERVES')
for i,(a,b,c) in enumerate([
 ('EVERYDAY USERS','Practice, learn and discover relevant programmes at their own pace.',GOLD),
 ('TEACHERS AND MENTORS','Share approved expertise and choose whether to offer further support.',PURPLE),
 ('SEEKERS AND MEMBERS','Ask for appropriate help and choose useful learning or mentoring referrals.',GREEN),
 ('VOLUNTEERS','Prepare for suitable duties and contribute within approved responsibilities.',GREEN),
 ('MRM OPERATIONS','Review applications, coordinate services and follow outcomes.',CYAN),
 ('PARTNERS AND SPONSORS','Support approved programmes through clear referral and recognition arrangements.',GOLD)]):
 card(s,.7+(i%3)*5,2.74+(i//3)*2.46,4.6,2.2,a,b,c)

s=base('Faith, learning and care strengthen each other','The connection is a relevant next step for a person—not simply another feature in a menu.','THE CONNECTED EXPERIENCE')
for i,(a,b,c) in enumerate([
 ('DAILY PRACTICE','A Quran reader asks for guidance and discovers a suitable beginner programme.',GOLD),
 ('APPROVED LEARNING','A teacher chooses to offer relevant support; a volunteer prepares for a duty.',PURPLE),
 ('COMMUNITY CARE','A seeker chooses a class; aggregate needs help MRM plan useful programmes.',GREEN)]):
 card(s,.7+i*5,2.94,4.6,3.05,a,b,c)
 if i<2:arrow(s,5.36+i*5,4.46,5.65+i*5,4.46)
card(s,.7,6.42,7.1,1.36,'VALUE FOR THE PERSON','A clearer next step, relevant expertise and less repeated explanation.',CYAN)
card(s,8.2,6.42,7.1,1.36,'VALUE FOR THE COMMUNITY','Better prepared contributors and more appropriate service handoffs.',GREEN)

s=base('Proposal roadmap','Explore the platform story, the user journeys, the feature detail and the delivery approach.','PRESENTATION GUIDE')
agenda=[]
for i,(a,b,c,target) in enumerate([
 ('01  Vision and community','The proposition, audiences and connected experience.',CYAN,2),
 ('02  Connected user journeys','Teacher, seeker and volunteer motivations and flows.',GREEN,6),
 ('03  Complete feature design','Everyday tools, learning, care and operating capabilities.',PURPLE,24),
 ('04  End-to-end services','Care, learning and shared service events.',GREEN,71),
 ('05  Extended concepts','Campaign, impact and contribution ideas for separate agreement.',GOLD,75),
 ('06  Delivery and next steps','Dependencies, responsibilities, validation and scope.',CYAN,80)]):
 x=.7+(i%3)*5;y=2.75+(i//3)*2.45;card(s,x,y,4.6,2.16,a,b,c);agenda.append((s.shapes[-2],target));agenda.append((s.shapes[-3],target))

s=base('A clear core scope, with deliberate extension decisions','Build the agreed service first, then assess additional concepts on their value and operating requirements.','SCOPE AND COMMERCIAL PRINCIPLES')
card(s,.7,2.8,7.1,2.16,'CORE SERVICE','Daily utilities, approved learning, volunteer support, controlled referrals, a public website and MRM administration.',CYAN)
card(s,8.2,2.8,7.1,2.16,'PARTICIPATION MODEL','Teacher-only subscription entitlements remain distinct from volunteer approval. Core member assistance and volunteering stay free.',GREEN)
card(s,.7,5.24,7.1,2.16,'EXTENSIONS FOR SEPARATE AGREEMENT','Campaign and contribution concepts require confirmed scope, payment arrangements, authority boundaries and operating ownership.',GOLD)
card(s,8.2,5.24,7.1,2.16,'IMPLEMENTATION FOUNDATION','Preserve the agreed Flutter / Zabi approach. Confirm integrations, acceptance criteria and service responsibilities before delivery.',PURPLE)

s=base('A connected platform MRM can launch and grow','Agree the priority journeys, confirm the operating model and turn the proposal into an executable delivery plan.','NEXT STEPS')
for i,(a,b,c) in enumerate([
 ('ALIGN THE EXPERIENCE','Validate the teacher, volunteer and seeker journeys with MRM and representative users.',GREEN),
 ('CONFIRM THE DELIVERY','Agree core scope, extension decisions, milestones, responsibilities and acceptance criteria.',CYAN),
 ('BUILD WITH CONFIDENCE','Deliver in dependency order, verify the real service flows and prepare the people who will operate them.',PURPLE)]):
 card(s,.7+i*5,2.91,4.6,3.22,a,b,c)
txt(s,.87,6.77,14.1,.9,'Everyday usefulness. Meaningful learning. Responsible care.',27,WHITE,True)
txt(s,.87,7.79,14.1,.3,'MRM Digital Community Platform  |  Zylax Technologies',13,CYAN)

# Import existing slide artwork without flattening editable content.
original=Presentation(ROOT.parent/'strategic-rebuild/MRM_Strategic_Platform_72_Pages.pptx')
bridges=Presentation(ROOT.parent/'strategic-rebuild/user-journeys/MRM_User_Value_Bridges_24_Pages.pptx')
newslides=list(P.slides)
# (source namespace, 1-based source page) is the stable cross-document key.
order=[('new',i) for i in range(1,6)]+[('bridge',20),('bridge',21)]+[('bridge',i) for i in range(5,20)]+[('original',i) for i in range(5,66)]+[('new',6),('new',7)]
sources={'new':newslides,'original':list(original.slides),'bridge':list(bridges.slides)}
dest={('new',i+1):s for i,s in enumerate(newslides)}
for group,page in order:
 if group!='new':dest[(group,page)]=P.slides.add_slide(P.slide_layouts[6])
# Destinations of removed editorial/reference slides route to the appropriate client section.
fallback={('original',1):('new',1),('original',2):('new',2),('original',3):('new',5),('original',4):('new',3),('original',66):('new',6),('original',67):('new',6),('original',68):('new',6),('original',69):('new',5),('original',70):('new',5),('original',71):('new',6),('original',72):('new',7),('bridge',1):('new',2),('bridge',2):('bridge',5),('bridge',3):('new',5),('bridge',4):('new',4),('bridge',22):('original',62),('bridge',23):('new',6),('bridge',24):('original',64)}
for group,page in order:
 if group=='new':continue
 src=sources[group][page-1];dst=dest[(group,page)]
 index={s.part.partname:i+1 for i,s in enumerate(sources[group])}
 for sh in src.shapes:
  el=deepcopy(sh.element)
  for node in el.iter():
   for attr in [qn('r:id'),qn('r:embed'),qn('r:link')]:
    rid=node.get(attr)
    if not rid:continue
    rel=src.part.rels[rid]
    if rel.reltype.endswith('/image'):_,nr=dst.part.get_or_add_image_part(io.BytesIO(rel.target_part.blob))
    elif rel.reltype.endswith('/slide'):
     key=(group,index[rel.target_part.partname]);key=fallback.get(key,key);nr=dst.part.relate_to(dest[key].part,rel.reltype)
    elif rel.is_external:nr=dst.part.relate_to(rel.target_ref,rel.reltype,is_external=True)
    else:raise ValueError(rel.reltype)
    node.set(attr,nr)
  dst.shapes._spTree.insert_element_before(el,'p:extLst')
 dst.notes_slide.notes_text_frame.text=src.notes_slide.notes_text_frame.text
# Set the final reading order, then replace all section-local numbering with a global page system.
by_id={s.slide_id:s for s in P.slides};idels={el.id:el for el in P.slides._sldIdLst}
ordered_ids=[dest[key].slide_id for key in order]
for el in list(P.slides._sldIdLst):P.slides._sldIdLst.remove(el)
for slide_id in ordered_ids:P.slides._sldIdLst.append(idels[slide_id])
assert len(P.slides)==85
# Remove inherited footers on all content slides; add a consistent, clickable client footer.
for i,s in enumerate(P.slides):
 if i==0:continue
 for sh in list(s.shapes):
  if sh.top>=Inches(8.3):s.shapes._spTree.remove(sh.element)
 txt(s,.7,8.49,4,.25,'MRM DIGITAL COMMUNITY PLATFORM',9,MUTED)
 s.shapes[-1].click_action.target_slide=dest[('new',5)]
 txt(s,5.0,8.49,7.8,.25,'PROJECT PROPOSAL  |  ZYLAX TECHNOLOGIES',8,MUTED,align=ns['PP_ALIGN'].CENTER)
 txt(s,14.1,8.43,1.2,.35,f'{i+1:02d} / 85',12,CYAN,True,ns['PP_ALIGN'].RIGHT)
# Professionalise source slide labels; source evidence qualifications remain in notes/body where useful.
replacements={
 'Care is a connected service, not a dispatch button':'From a support request to accountable care',
 'Why every capability belongs in the overall product':'Platform priorities and delivery dependencies',
 'The overall network: features hand off real work':'The shared service foundation',
 'PROPOSED USER JOURNEY':'CONNECTED USER JOURNEY'}
for s in P.slides:
 for sh in s.shapes:
  if not sh.has_text_frame:continue
  for p in sh.text_frame.paragraphs:
   for r in p.runs:
    for a,b in replacements.items():r.text=r.text.replace(a,b)
# Native agenda links use the global reading order.
for sh,page in agenda:sh.click_action.target_slide=P.slides[page-1]
# Chapter-level identifiers for bookmarks and audit traceability.
oldtitles=json.loads((ROOT.parent/'strategic-rebuild/Main_Contents.json').read_text());btitles=json.loads((ROOT.parent/'strategic-rebuild/user-journeys/Contents.json').read_text())
titles=[]
for group,page in order:
 t=ns['TITLES'][page-1] if group=='new' else (oldtitles if group=='original' else btitles)[page-1]
 for a,b in replacements.items():t=t.replace(a,b)
 titles.append(t)
P.core_properties.title='MRM Digital Community Platform — Project Proposal'
P.core_properties.subject='Complete platform, connected user journeys, feature design and delivery'
P.core_properties.author='Zylax Technologies';P.core_properties.keywords='MRM, Muslim community, learning, volunteer support, platform proposal'
P.save(ROOT/'MRM_Complete_Client_Proposal.pptx')
(ROOT/'Contents.json').write_text(json.dumps(titles,indent=2))
(ROOT/'Assembly_Model.json').write_text(json.dumps([{'page':i+1,'source':g,'source_page':p,'title':titles[i]} for i,(g,p) in enumerate(order)],indent=2))
print('Built complete 85-page client proposal with a three-screen app cover.')
